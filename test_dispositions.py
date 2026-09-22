"""Final QuickBooks dispositions and the proposed-JE population.

Every QuickBooks source row ends in exactly one of MATCHED,
EXACT_QBO_DUPLICATE_EXCLUDED, REVIEW_HOLD, or TRUE_UNMATCHED; only TRUE_UNMATCHED
rows (plus the documented PO Re-use Error groups) feed the proposed journal entry.
"""

import copy
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from fuzzy_po_matching import is_fuzzy_po_match, significant_po_tokens
from matching import (
    DISPOSITION_DUPLICATE_EXCLUDED,
    DISPOSITION_MATCHED,
    DISPOSITION_REVIEW_HOLD,
    DISPOSITION_TRUE_UNMATCHED,
    FINAL_DISPOSITIONS,
    QB_ID,
    build_reconciliation,
    validate_reconciliation,
)
from workpapers import build_analytics_workbook, build_primary_workbook


def _qb(po, invoice, amount, period="1", **extra):
    return {"PO": po, "Invoice": invoice, "Amount": amount, "Qty": 1, "Period": period,
            "Customer": "Acme", "Date": "2026-01-05", **extra}


def _inf(po, invoice, amount, period="1"):
    return {"PO": po, "Invoice": invoice, "Amount": amount, "Period": period,
            "Customer": "Acme", "Date": "2026-01-05"}


def _reconcile(qb_rows, inf_rows, qb_mapping, inf_mapping, make_metadata, **kwargs):
    return build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping, make_metadata(), 2026, **kwargs,
    )


def _ledger(result):
    return result.qb_dispositions.set_index("QBO Row ID")


def _by_disposition(result, disposition):
    frame = result.qb_dispositions
    return list(frame.loc[frame["Final Disposition"] == disposition, "QBO Row ID"])


@pytest.fixture
def mixed_result(qb_mapping, inf_mapping, make_metadata):
    """One row of every outcome: matched, grouped match, exact duplicate, a PO
    already represented, an amount conflict on a shared invoice, and a genuine
    missing transaction."""
    qb_rows = [
        _qb("PO1", "INV1", 100.00),                     # QB-1 matched
        _qb("PO1", "INV1", 100.00),                     # QB-2 exact duplicate of QB-1
        _qb("PO-G", "INV-A", 60.00),                    # QB-3 + QB-4 grouped match
        _qb("PO-G", "INV-B", 40.00),
        _qb("PO1", "INV1B", 5.00),                      # QB-5 PO already represented by M-001
        _qb("PO-Y", "INV-V", 100.00),                   # QB-6 / QB-7 share an invoice with an INF row at 90
        _qb("PO-Z", "INV-V", 100.00),
        _qb("PO-NONE", "INV-NONE", 33.00),              # QB-8 genuinely missing from Infinium
    ]
    inf_rows = [
        _inf("PO1", "INV1", 100.00),
        _inf("PO-G", "", 100.00),
        _inf("PO-Q", "INV-V", 90.00),
    ]
    return _reconcile(qb_rows, inf_rows, qb_mapping, inf_mapping, make_metadata)


# ---------------------------------------------------------------------------
# 1. Every source row keeps exactly one disposition, and it all ties out
# ---------------------------------------------------------------------------

def test_every_source_row_has_exactly_one_disposition_and_totals_tie(mixed_result):
    result = mixed_result
    ledger = result.qb_dispositions
    assert list(ledger["QBO Row ID"]) == [f"QB-{n}" for n in range(1, 9)]      # every row, in source order
    assert set(ledger["Final Disposition"]) <= set(FINAL_DISPOSITIONS)
    counts = ledger["Final Disposition"].value_counts().to_dict()
    assert sum(counts.values()) == len(result.qb_work) == result.metrics["QuickBooks Rows"]
    assert round(ledger["Amount"].sum(), 2) == pytest.approx(result.metrics["QuickBooks Source Total"])
    assert result.metrics["Control Status"] == "PASS"

    checks = result.controls.set_index("Check")
    for name in (
        "Source QuickBooks rows = Matched + Confirmed Duplicate Excluded + Review-Hold + True-Unmatched rows",
        "Every QuickBooks source row has exactly one final disposition",
        "Source QuickBooks amount = Matched + Confirmed Duplicate Excluded + Review-Hold + True-Unmatched amounts",
        "Proposed JE = sum of TRUE_UNMATCHED QuickBooks amounts",
    ):
        assert checks.loc[name, "Status"] == "PASS", name
    assert result.metrics["Final Disposition - Matched Rows"] == counts[DISPOSITION_MATCHED]


def test_each_row_lands_in_the_expected_disposition_with_a_precise_reason(mixed_result):
    ledger = _ledger(mixed_result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED
    assert ledger.loc["QB-1", "Final Reason"] == "Exact PO + Invoice + Amount"
    assert ledger.loc["QB-1", "Match Ref."] == "M-001"

    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_DUPLICATE_EXCLUDED
    assert ledger.loc["QB-2", "Final Reason"].startswith("Confirmed Exact QBO Duplicate — Excess Copy Excluded")
    assert "confirmed by stable-field fingerprint" in ledger.loc["QB-2", "Final Reason"]
    assert ledger.loc["QB-2", "Canonical QBO Row ID"] == "QB-1"
    assert ledger.loc["QB-2", "Duplicate Group ID"].startswith("DUP-")

    for row_id in ("QB-3", "QB-4"):                       # aggregate: every contributing row kept
        assert ledger.loc[row_id, "Final Disposition"] == DISPOSITION_MATCHED
        assert ledger.loc[row_id, "Aggregate Group ID"] == "G-001"
        assert ledger.loc[row_id, "Final Reason"] == "Aggregate PO + Exact Amount"

    assert ledger.loc["QB-5", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-5", "Reason Code"] == "REVIEW_HOLD_PO_ALREADY_REPRESENTED"
    assert ledger.loc["QB-5", "Final Reason"] == "Review Hold — PO Already Represented by Match M-001"
    assert ledger.loc["QB-5", "Related Match Ref."] == "M-001"
    assert ledger.loc["QB-5", "Review ID"].startswith("REV-")
    assert ledger.loc["QB-5", "In Proposed JE"] == "No"

    assert ledger.loc["QB-8", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED
    assert ledger.loc["QB-8", "Final Reason"] == "True Unmatched — No Remaining Infinium Candidate"
    assert ledger.loc["QB-8", "In Proposed JE"] == "Yes"


def test_the_proposed_je_is_the_sum_of_true_unmatched_rows_only(mixed_result):
    result = mixed_result
    ledger = result.qb_dispositions
    true_unmatched = ledger.loc[ledger["Final Disposition"] == DISPOSITION_TRUE_UNMATCHED, "Amount"].sum()
    assert result.metrics["Proposed JE Amount"] == pytest.approx(true_unmatched)
    assert result.metrics["Unresolved QuickBooks Amount"] == pytest.approx(true_unmatched)
    # Nothing that is a duplicate, a hold, or a match feeds it.
    fed = set(ledger.loc[ledger["In Proposed JE"] == "Yes", "Final Disposition"])
    assert fed == {DISPOSITION_TRUE_UNMATCHED}
    assert not set(ledger.loc[ledger["In Proposed JE"] == "Yes", "Reason Code"]) & {
        code for code in ledger["Reason Code"] if code.startswith("REVIEW_HOLD")
    }


def test_validation_rejects_a_ledger_that_disagrees_with_the_reconciliation(mixed_result):
    validate_reconciliation(mixed_result)

    lost = copy.deepcopy(mixed_result)
    lost.qb_dispositions = lost.qb_dispositions.iloc[:-1].reset_index(drop=True)
    with pytest.raises(ValueError, match="exactly one final disposition"):
        validate_reconciliation(lost)

    leaked = copy.deepcopy(mixed_result)
    row = leaked.qb_dispositions["QBO Row ID"] == "QB-5"       # a review hold sneaking into the JE
    leaked.qb_dispositions.loc[row, "Final Disposition"] = DISPOSITION_TRUE_UNMATCHED
    with pytest.raises(ValueError, match="do not agree with the reconciliation populations"):
        validate_reconciliation(leaked)

    forced = copy.deepcopy(mixed_result)
    forced.qb_dispositions.loc[forced.qb_dispositions["QBO Row ID"] == "QB-5", "In Proposed JE"] = "Yes"
    with pytest.raises(ValueError, match="feeds the journal entry"):
        validate_reconciliation(forced)


# ---------------------------------------------------------------------------
# 2. Exact QuickBooks duplicates are decided before matching
# ---------------------------------------------------------------------------

def test_an_exact_duplicate_is_excluded_before_matching_and_never_reaches_the_je(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_qb("PO1", "INV1", 100.00), _qb("PO1", "INV1", 100.00), _qb("PO1", "INV1", 100.00)],
        [_inf("PO1", "INV1", 100.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert _by_disposition(result, DISPOSITION_MATCHED) == ["QB-1"]                 # the canonical row
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2", "QB-3"]
    assert result.unmatched_qb == [] and result.metrics["Proposed JE Amount"] == 0
    ledger = _ledger(result)
    assert set(ledger.loc[["QB-2", "QB-3"], "Canonical QBO Row ID"]) == {"QB-1"}
    assert ledger.loc["QB-2", "Duplicate Group ID"] == ledger.loc["QB-3", "Duplicate Group ID"] != ""
    # The excess copies never entered a matching pass.
    assert all(1 not in group.qb_rows and 2 not in group.qb_rows for group in result.matches)


def test_the_canonical_row_is_the_earliest_by_source_order(qb_mapping, inf_mapping, make_metadata):
    result = _reconcile(
        [_qb("PO9", "INV9", 12.00), _qb("PO1", "INV1", 100.00), _qb("PO1", "INV1", 100.00)],
        [_inf("PO-X", "INV-X", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-3", "Canonical QBO Row ID"] == "QB-2"
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED       # canonical stays a real exception
    assert ledger.loc["QB-3", "Final Disposition"] == DISPOSITION_DUPLICATE_EXCLUDED


def test_duplicate_identity_is_normalized_and_uses_signed_cents(qb_mapping, inf_mapping, make_metadata):
    result = _reconcile(
        [
            _qb(" po-1 ", "inv 1", 100.10),         # QB-1
            _qb("PO-1", "INV1", 100.10),            # QB-2: same identity after normalization
            _qb("PO-1", "INV1", -100.10),           # QB-3: the reversal is NOT a duplicate
            _qb("PO-1", "INV1", 100.1),             # QB-4: same signed cents as QB-1
        ],
        [_inf("PO-X", "INV-X", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert list(ledger.loc[ledger["Final Disposition"] == DISPOSITION_DUPLICATE_EXCLUDED].index) == ["QB-2", "QB-4"]
    assert ledger.loc["QB-3", "Final Disposition"] != DISPOSITION_DUPLICATE_EXCLUDED
    assert ledger.loc["QB-1", "Final Disposition"] != DISPOSITION_DUPLICATE_EXCLUDED


def test_a_weak_basis_sibling_of_a_matched_row_is_held_not_accrued(qb_mapping, inf_mapping, make_metadata):
    """Only a PO (no invoice) ties these two QuickBooks rows together. One matches;
    the other is already-represented evidence -- held for review, not accrued."""
    result = _reconcile(
        [_qb("PO-W", "", 60.00), _qb("PO-W", "", 60.00)],
        [_inf("PO-W", "", 60.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"
    assert ledger.loc["QB-2", "Related Match Ref."] == "M-001"
    assert result.metrics["Proposed JE Amount"] == 0


# ---------------------------------------------------------------------------
# 3. Consumed candidates leave the pool
# ---------------------------------------------------------------------------

def test_a_consumed_candidate_no_longer_makes_a_later_candidate_ambiguous(
    qb_mapping, inf_mapping, make_metadata,
):
    """Two rows share a PO and an amount on each side. The pass on PO + Invoice +
    Amount consumes one pair; the other pair must then match on PO + Amount, not be
    rejected as ambiguous because of a record that is already spoken for."""
    result = _reconcile(
        [_qb("PO-X", "INV-1", 100.00), _qb("PO-X", "INV-2", 100.00)],
        [_inf("PO-X", "INV-1", 100.00), _inf("PO-X", "INV-9", 100.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    methods = sorted(group.method for group in result.matches)
    assert methods == ["PO + Amount", "PO + Invoice + Amount"]
    assert result.unmatched_qb == [] and result.unmatched_inf == []


def test_a_consumed_historical_record_no_longer_makes_a_later_one_ambiguous(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_qb("PO-X", "INV-1", 100.00), _qb("PO-X", "INV-2", 100.00)],
        [_inf("PO-ZZZ", "INV-ZZZ", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
        inf_secondary_raw=pd.DataFrame([_inf("PO-X", "INV-1", 100.00, "12"), _inf("PO-X", "INV-3", 100.00, "12")]),
        inf_secondary_mapping=inf_mapping,
    )
    ledger = _ledger(result)
    assert set(ledger["Final Disposition"]) == {DISPOSITION_MATCHED}
    assert len(result.historical_clearances["Clearance ID"].unique()) == 2


# ---------------------------------------------------------------------------
# 4/5. Aggregates keep every row; fuzzy PO is contained, symmetric, and held
# ---------------------------------------------------------------------------

def test_an_aggregate_match_keeps_every_contributing_row_under_one_group_id(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_qb("PO-G", "INV-A", 30.00), _qb("PO-G", "INV-B", 70.00)],
        [_inf("PO-G", "", 100.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = result.qb_dispositions
    assert list(ledger["QBO Row ID"]) == ["QB-1", "QB-2"]
    assert set(ledger["Aggregate Group ID"]) == {"G-001"} and set(ledger["Match Ref."]) == {"G-001"}
    register = result.match_register.iloc[0]
    assert register["QuickBooks Amount"] == register["Infinium Amount"] == 100.00
    assert register["Amount Difference"] == 0


@pytest.mark.parametrize("qb_po, inf_po", [
    ("PRIME STAINLESS - TIM FERRIS", "PO-TIM FERRIS"),
    ("PO-TIM FERRIS", "PRIME STAINLESS - TIM FERRIS"),        # containment is symmetric
])
def test_fuzzy_po_uses_whole_word_containment_in_either_direction(qb_po, inf_po):
    assert is_fuzzy_po_match(significant_po_tokens(qb_po), significant_po_tokens(inf_po))


def test_fuzzy_po_requires_the_whole_shorter_reference_not_a_partial_overlap():
    assert not is_fuzzy_po_match(significant_po_tokens("TIM SMITH"), significant_po_tokens("TIM FERRIS"))
    assert not is_fuzzy_po_match(significant_po_tokens("TIM"), significant_po_tokens("PO-TIM FERRIS"))   # too generic alone


def test_a_fuzzy_match_needs_the_exact_amount_stays_fuzzy_and_never_feeds_the_je(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [
            _qb("PRIME STAINLESS - TIM FERRIS", "Q-1", 250.00),        # fuzzy candidate, exact amount
            _qb("PRIME STAINLESS - TIM FERRIS", "Q-2", 999.00),        # same words, different amount
        ],
        [_inf("PO-TIM FERRIS", "I-1", 250.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-1", "Reason Code"] == "REVIEW_HOLD_FUZZY_MATCH"
    assert "Fuzzy PO" in ledger.loc["QB-1", "Final Reason"] and "Exact Amount" in ledger.loc["QB-1", "Final Reason"]
    assert ledger.loc["QB-1", "In Proposed JE"] == "No"
    assert ledger.loc["QB-2", "Final Disposition"] != DISPOSITION_MATCHED           # amount must agree exactly
    assert result.matches == []                                                    # never an accepted match


# ---------------------------------------------------------------------------
# 6-9. "Cannot safely match" is not "does not exist in Infinium"
# ---------------------------------------------------------------------------

def test_reference_evidence_holds_carry_their_full_audit_detail(mixed_result):
    hold = mixed_result.reference_hold_analysis.set_index("QuickBooks Row ID").loc["QB-5"]
    assert hold["Reason Code"] == "REVIEW_HOLD_PO_ALREADY_REPRESENTED"
    assert hold["Related Match Ref."] == "M-001"
    assert hold["Related QuickBooks Row IDs"] == "QB-1"
    assert hold["Related Infinium Row IDs"] == "INF-1"
    assert hold["QuickBooks Amount"] == 5.00
    assert hold["Posting Disposition"] == "REVIEW REQUIRED - DO NOT POST"
    assert "Hold" in hold["Hold ID"] or hold["Hold ID"].startswith("RHOLD-")


def test_an_amount_variance_against_a_shared_invoice_is_a_hold_not_an_accrual(mixed_result):
    ledger = _ledger(mixed_result)
    for row_id in ("QB-6", "QB-7"):
        assert ledger.loc[row_id, "Final Disposition"] == DISPOSITION_REVIEW_HOLD
        assert ledger.loc[row_id, "Reason Code"] == "REVIEW_HOLD_AMOUNT_VARIANCE"
        assert ledger.loc[row_id, "Final Reason"] == "Review Hold — Amount Variance"
    hold = mixed_result.reference_hold_analysis.set_index("QuickBooks Row ID").loc["QB-6"]
    assert hold["QuickBooks Amount"] == 100.00 and hold["Infinium Amount"] == 90.00
    assert hold["Amount Difference"] == 10.00 and hold["Related Infinium Row IDs"] == "INF-3"


def test_an_exact_candidate_that_is_not_uniquely_available_is_held(qb_mapping, inf_mapping, make_metadata):
    result = _reconcile(
        [_qb("PO-X", "INV-1", 100.00), _qb("PO-Y", "INV-1", 100.00)],     # same invoice, same amount
        [_inf("PO-Z", "INV-1", 100.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert set(ledger["Reason Code"]) == {"REVIEW_HOLD_EXACT_CANDIDATE_NOT_UNIQUE"}
    assert result.metrics["Proposed JE Amount"] == 0


def test_an_unreadable_amount_is_held_rather_than_treated_as_a_missing_transaction(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_qb("PO-BAD", "INV-BAD", "not a number"), _qb("PO-N", "INV-N", 10.00)],
        [_inf("PO-Z", "INV-Z", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Reason Code"] == "REVIEW_HOLD_INVALID_AMOUNT"
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED
    assert result.metrics["Proposed JE Amount"] == pytest.approx(10.00)


def test_duplicate_and_hold_wording_never_appears_in_the_accrual_population(mixed_result):
    """No row whose status speaks of a duplicate candidate, a PO already used, an
    already-represented record, or an amount variance may sit in the JE table."""
    banned = ("duplicate", "already used", "already represented", "amount variance", "typo")
    for row in mixed_result.paired_rows:
        if row["Section"] == "02 Unmatched QuickBooks":
            text = f"{row['Match Result']} {row['Exception Cause']}".lower()
            assert not any(word in text for word in banned), text


def test_infinium_only_exceptions_never_offset_the_quickbooks_je(qb_mapping, inf_mapping, make_metadata):
    base = _reconcile(
        [_qb("PO-NONE", "INV-NONE", 500.00)], [_inf("PO-Z", "INV-Z", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    with_credit = _reconcile(
        [_qb("PO-NONE", "INV-NONE", 500.00)],
        [_inf("PO-Z", "INV-Z", 1.00), _inf("BILLING", "DUPLICATE BILLING", -500.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert base.metrics["Proposed JE Amount"] == with_credit.metrics["Proposed JE Amount"] == 500.00
    assert with_credit.metrics["Unmatched Infinium Rows"] == 2          # still visible, in its own total
    assert with_credit.metrics["Unmatched Infinium Amount"] == pytest.approx(-499.00)


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def test_the_analytics_workbook_lists_every_source_row_in_the_disposition_ledger(mixed_result):
    wb = load_workbook(io.BytesIO(build_analytics_workbook(mixed_result)))
    ws = wb["QB Disposition Ledger"]
    header_row = next(row for row in ws.iter_rows(max_row=8) if any(c.value == "Final Disposition" for c in row))
    headers = [c.value for c in header_row]
    rows = [
        dict(zip(headers, [c.value for c in row]))
        for row in ws.iter_rows(min_row=header_row[0].row + 1) if row[0].value
    ]
    assert [r["QBO Row ID"] for r in rows] == list(mixed_result.qb_dispositions["QBO Row ID"])
    assert {r["Final Disposition"] for r in rows} <= set(FINAL_DISPOSITIONS)
    text = " ".join(str(c.value) for row in wb["Executive Summary"].iter_rows() for c in row if c.value)
    assert "Proposed JE = sum of TRUE_UNMATCHED QuickBooks amounts" in text


def test_the_exceptions_table_holds_only_true_unmatched_rows_and_holds_are_itemized_below(mixed_result):
    ws = load_workbook(io.BytesIO(build_primary_workbook(mixed_result)))["Unresolved Exceptions"]
    header_row = next(row for row in ws.iter_rows() if any(c.value == "Exception Status" for c in row))
    headers = [c.value for c in header_row]
    status_col = headers.index("Exception Status") + 1
    statuses = [
        ws.cell(r, status_col).value
        for r in range(header_row[0].row + 1, header_row[0].row + 1 + len(mixed_result.unmatched_qb))
    ]
    assert statuses == ["No matching Infinium records"]
    # The held rows are itemized in their own section, with the reason for each.
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "REVIEW HOLD | REFERENCE EVIDENCE IN INFINIUM - NOT ACCRUED" in text
    assert "Review Hold — PO Already Represented by Match M-001" in text
    assert "Review Hold — Amount Variance" in text


# ---------------------------------------------------------------------------
# Refinement 1: a shared key identifies a duplicate RELATIONSHIP; only a
# confirmed copy of one underlying transaction is excluded
# ---------------------------------------------------------------------------

def _stable(qb_mapping):
    """A mapping that also names the customer and transaction-date columns."""
    return {**qb_mapping, "customer": "Customer", "date": "Date"}


def _sales(customer, day, **extra):
    return _qb("PO1", "INV1", 100.00, Customer=customer, Date=day, **extra)


def test_two_legitimate_sales_with_the_same_key_but_different_identities_are_not_auto_excluded(
    qb_mapping, inf_mapping, make_metadata,
):
    """Same PO, invoice, and amount -- but a different customer and date: two real
    sales, not one entered twice. The sibling of the matched row is held for review
    (never discarded, never accrued) and points at the match."""
    result = _reconcile(
        [_sales("Acme", "2026-01-05"), _sales("Beta", "2026-01-09")],
        [_inf("PO1", "INV1", 100.00)],
        _stable(qb_mapping), inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"
    assert ledger.loc["QB-2", "Related Match Ref."] == "M-001"
    assert ledger.loc["QB-2", "Canonical QBO Row ID"] == "QB-1"
    assert "Potential Duplicate of Match M-001" in ledger.loc["QB-2", "Final Reason"]
    assert ledger.loc["QB-2", "In Proposed JE"] == "No"
    assert result.duplicate_qb_rows == [] and result.metrics["Proposed JE Amount"] == 0
    # The evidence for why it was not confirmed stays visible.
    report = result.duplicate_analysis.set_index("Source Row ID")
    assert not report.loc["QB-2", "Payload Confirmed"]
    assert {"Customer", "Date"} <= set(report.loc["QB-2", "Differing Confirmation Fields"].split("; "))


def test_the_representative_of_an_unconfirmed_group_stays_a_true_exception_and_the_rest_are_held(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_sales("Acme", "2026-01-05"), _sales("Beta", "2026-01-09"), _sales("Cato", "2026-01-11")],
        [_inf("PO-X", "INV-X", 1.00)],
        _stable(qb_mapping), inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED
    for row_id in ("QB-2", "QB-3"):
        assert ledger.loc[row_id, "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"
        assert ledger.loc[row_id, "Canonical QBO Row ID"] == "QB-1"
        assert "suspected canonical row" in ledger.loc[row_id, "Final Reason"]
    assert result.metrics["Proposed JE Amount"] == pytest.approx(100.00)        # only the representative accrues
    assert set(ledger.index) == {"QB-1", "QB-2", "QB-3"}                         # nothing disappears


def test_an_actual_copied_transaction_is_a_confirmed_duplicate(qb_mapping, inf_mapping, make_metadata):
    """The same record entered twice -- identical in every source column -- is a
    confirmed exact duplicate: one canonical row, the copy excluded and visible."""
    result = _reconcile(
        [_sales("Acme", "2026-01-05"), _sales("Acme", "2026-01-05")],
        [_inf("PO1", "INV1", 100.00)],
        _stable(qb_mapping), inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_DUPLICATE_EXCLUDED
    assert ledger.loc["QB-2", "Reason Code"] == "EXACT_QBO_DUPLICATE_EXCESS_COPY"
    assert ledger.loc["QB-2", "Canonical QBO Row ID"] == "QB-1"
    report = result.duplicate_analysis.set_index("Source Row ID")
    assert report.loc["QB-2", "Confirmation Basis"] == "Stable-field fingerprint (Customer, Date, Qty)"
    assert result.metrics["Proposed JE Amount"] == 0


def test_a_formatting_difference_in_the_key_columns_does_not_block_confirmation(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_qb(" po-1 ", "inv 1", 100.10, Customer="Acme"), _qb("PO-1", "INV1", 100.1, Customer="Acme")],
        [_inf("PO-X", "INV-X", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]


def test_only_a_line_level_id_can_independently_confirm_a_copy(qb_mapping, inf_mapping, make_metadata):
    """A mapped line-level source ID confirms a copied row even when the mapped
    fields alone would be too thin to; one that is not line-level cannot."""
    thin = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "quantity": "Qty", "period": None,
            "product": None, "line_id": "LineID"}
    base = dict(PO="PO1", Invoice="INV1", Amount=100.00, Qty=1)

    def run(mapping, rows):
        return _reconcile(rows, [_inf("PO-X", "INV-X", 1.00)], mapping, inf_mapping, make_metadata)

    copy = run(thin, [dict(base, LineID="L-100"), dict(base, LineID="L-100")])
    assert _by_disposition(copy, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]
    assert copy.duplicate_analysis.set_index("Source Row ID").loc["QB-2", "Confirmation Basis"] == "Line-level source ID"
    assert "confirmed by line-level source id" in _ledger(copy).loc["QB-2", "Final Reason"]

    # The same column, but shared by two DIFFERENT lines of one invoice: it is
    # invoice-level, so the ID is not trusted and confirms nothing.
    mapping = {**thin, "product": "Item"}
    invoice_level = run(mapping, [dict(base, Item="Widget", LineID="INV1"), dict(base, Item="Gadget", LineID="INV1")])
    assert _by_disposition(invoice_level, DISPOSITION_DUPLICATE_EXCLUDED) == []
    assert _ledger(invoice_level).loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"


def test_a_transaction_level_id_is_supporting_evidence_only(qb_mapping, inf_mapping, make_metadata):
    mapping = {**_stable(qb_mapping), "transaction_id": "TxnID"}
    base = dict(PO="PO1", Invoice="INV1", Amount=100.00, Qty=1, Period="1", Customer="Acme", Date="2026-01-05")

    def run(first, second, map_=mapping):
        return _reconcile([first, second], [_inf("PO-X", "INV-X", 1.00)], map_, inf_mapping, make_metadata)

    # Same transaction ID AND an identical, sufficient line fingerprint => confirmed.
    assert _by_disposition(run(dict(base, TxnID="T-1"), dict(base, TxnID="T-1")), DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]
    # Same transaction ID but the line differs => not a copy.
    assert _by_disposition(
        run(dict(base, TxnID="T-1"), dict(base, TxnID="T-1", Customer="Beta")), DISPOSITION_DUPLICATE_EXCLUDED,
    ) == []
    # Different transaction IDs => different transactions.
    assert _by_disposition(run(dict(base, TxnID="T-1"), dict(base, TxnID="T-2")), DISPOSITION_DUPLICATE_EXCLUDED) == []
    # A shared transaction ID with only weak fields is not enough.
    weak = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "quantity": "Qty", "period": "Period",
            "product": None, "transaction_id": "TxnID"}
    result = run(dict(base, TxnID="T-1"), dict(base, TxnID="T-1"), weak)
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == []
    assert _ledger(result).loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"


def test_fiscal_period_is_context_not_identity(qb_mapping, inf_mapping, make_metadata):
    """The same sale appearing under a different fiscal period is still a copy."""
    rows = [_sales("Acme", "2026-01-05", Period="5"), _sales("Acme", "2026-01-05", Period="6")]
    result = _reconcile(rows, [_inf("PO-X", "INV-X", 1.00)], _stable(qb_mapping), inf_mapping, make_metadata)
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]
    assert "Period" not in result.duplicate_analysis.set_index("Source Row ID").loc["QB-2", "Confirmation Basis"]


def test_quantity_and_period_alone_never_confirm_an_exclusion(qb_mapping, inf_mapping, make_metadata):
    weak = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "quantity": "Qty", "period": "Period", "product": None}
    result = _reconcile(
        [_qb("PO1", "INV1", 100.00), _qb("PO1", "INV1", 100.00)], [_inf("PO-X", "INV-X", 1.00)],
        weak, inf_mapping, make_metadata,
    )
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == []
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED       # the representative
    assert ledger.loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"      # never excluded, never accrued
    assert result.metrics["Proposed JE Amount"] == pytest.approx(100.00)
    reason = result.duplicate_analysis.set_index("Source Row ID").loc["QB-2", "Potential Duplicate Reason"]
    assert "not sufficient" in reason


def test_line_level_attributes_item_and_rate_strengthen_and_distinguish(qb_mapping, inf_mapping, make_metadata):
    mapping = {**_stable(qb_mapping), "product": "Item", "rate": "Rate"}
    a = _sales("Acme", "2026-01-05", Item="Widget", Rate=25.0)
    same = _reconcile([a, dict(a)], [_inf("PO-X", "INV-X", 1.00)], mapping, inf_mapping, make_metadata)
    assert _by_disposition(same, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]
    basis = same.duplicate_analysis.set_index("Source Row ID").loc["QB-2", "Confirmation Basis"]
    assert basis == "Stable-field fingerprint (Customer, Date, Item, Qty, Rate)"
    for change in ({"Item": "Gadget"}, {"Rate": 30.0}):
        result = _reconcile([a, {**a, **change}], [_inf("PO-X", "INV-X", 1.00)], mapping, inf_mapping, make_metadata)
        assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == [], change


# ---------------------------------------------------------------------------
# Refinement 2: PO re-use and "already represented" are holds
# ---------------------------------------------------------------------------

def test_po_reuse_rows_with_infinium_evidence_are_held_and_without_it_stay_true_unmatched(
    qb_mapping, inf_mapping, make_metadata,
):
    held = _reconcile(
        [_qb("PO-R", "INV-A", 100.00), _qb("PO-R", "INV-B", 50.00)],
        [_inf("PO-R", "", 140.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(held)
    assert set(ledger["Reason Code"]) == {"REVIEW_HOLD_PO_REUSE"}
    assert set(ledger["Related Infinium Row IDs"]) == {"INF-1"}
    assert held.metrics["Proposed JE Amount"] == 0 and len(held.po_reuse_errors) == 1

    unsupported = _reconcile(
        [_qb("PO-R", "INV-A", 100.00), _qb("PO-R", "INV-B", 50.00)],
        [_inf("PO-OTHER", "INV-OTHER", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert set(_ledger(unsupported)["Final Disposition"]) == {DISPOSITION_TRUE_UNMATCHED}
    assert unsupported.metrics["Proposed JE Amount"] == pytest.approx(150.00)


# ---------------------------------------------------------------------------
# Refinement 3: controlled typo tolerance
# ---------------------------------------------------------------------------

def test_a_controlled_typo_matches_when_the_amount_is_exact_and_is_labeled_as_not_exact(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [_qb("ELIOT ELECTRIC", "Q-1", 300.00)], [_inf("ELLIOT ELECTRIC", "I-1", 300.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    (group,) = result.matches
    assert group.method == "Controlled PO Typo + Exact Amount"
    assert "'ELIOT' (QuickBooks) vs 'ELLIOT' (Infinium)" in group.explanation
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED
    assert ledger.loc["QB-1", "Reason Code"] == "MATCH_CONTROLLED_PO_TYPO"
    assert ledger.loc["QB-1", "Final Reason"] == "Controlled PO Typo + Exact Amount"
    assert ledger.loc["QB-1", "Match Ref."] == "M-001"
    assert result.metrics["Proposed JE Amount"] == 0


def test_a_controlled_typo_never_matches_on_a_different_amount(qb_mapping, inf_mapping, make_metadata):
    result = _reconcile(
        [_qb("ELIOT ELECTRIC", "Q-1", 300.00)], [_inf("ELLIOT ELECTRIC", "I-1", 300.01)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert result.matches == []
    assert _ledger(result).loc["QB-1", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED


def test_the_typo_rule_is_deterministic_not_a_similarity_score():
    from fuzzy_po_matching import controlled_typo_pair

    tokens = significant_po_tokens
    assert controlled_typo_pair(tokens("ELIOT ELECTRIC"), tokens("ELLIOT ELECTRIC")) == ("ELIOT", "ELLIOT")   # insertion
    assert controlled_typo_pair(tokens("ELLIOT ELECTRIC"), tokens("ELIOT ELECTRIC")) == ("ELLIOT", "ELIOT")   # deletion
    assert controlled_typo_pair(tokens("HARPER SUPPLY"), tokens("HARPOR SUPPLY")) == ("HARPER", "HARPOR")     # substitution, strong context
    assert controlled_typo_pair(tokens("HOPPER"), tokens("HOOPER")) is None          # lone substitution: a different name
    assert controlled_typo_pair(tokens("HOPPER"), tokens("SHOPPER")) is None         # different first letter
    assert controlled_typo_pair(tokens("ELIOT ELECTRIC"), tokens("ELLIOTT ELECTRICAL")) is None   # two differences
    assert controlled_typo_pair(tokens("BOLTS"), tokens("BOLT")) is None             # too short to be a meaningful token


def test_two_possible_typo_candidates_are_sent_to_review_not_matched(qb_mapping, inf_mapping, make_metadata):
    result = _reconcile(
        [_qb("ELIOT ELECTRIC", "Q-1", 300.00)],
        [_inf("ELLIOT ELECTRIC", "I-1", 300.00), _inf("ELLIOT ELECTRIC", "I-2", 300.00)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert result.matches == []                                          # nothing guessed
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-1", "Reason Code"] == "REVIEW_HOLD_TYPO_CANDIDATES"
    assert ledger.loc["QB-1", "Related Infinium Row IDs"] == "INF-1; INF-2"
    assert result.metrics["Proposed JE Amount"] == 0


# ---------------------------------------------------------------------------
# Refinement 4: apparent support that is withheld by a historical issue
# ---------------------------------------------------------------------------

def test_a_row_supported_only_by_a_withheld_historical_record_is_held(qb_mapping, inf_mapping, make_metadata):
    result = _reconcile(
        [_qb("PO-H", "INV-H", 100.00), _qb("PO-NONE", "INV-NONE", 7.00)],
        [_inf("PO-ZZZ", "INV-ZZZ", 1.00)],
        qb_mapping, inf_mapping, make_metadata,
        # Two historical rows share PO and amount with no invoice: an unresolved
        # historical-duplicate issue withholds BOTH from clearing anything.
        inf_secondary_raw=pd.DataFrame([_inf("PO-H", "", 100.00, "12"), _inf("PO-H", "", 100.00, "12")]),
        inf_secondary_mapping=inf_mapping,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-1", "Reason Code"] == "REVIEW_HOLD_HISTORICAL_CLEARANCE"
    assert ledger.loc["QB-1", "Related Infinium Row IDs"] == "INF-HIST-1; INF-HIST-2"
    hold = result.reference_hold_analysis.set_index("QuickBooks Row ID").loc["QB-1"]
    assert "DUP-" in hold["Explanation"]                                  # names the blocking duplicate group
    # A row with no support anywhere, historical or current, is still a genuine unmatched item.
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED
    assert result.metrics["Proposed JE Amount"] == pytest.approx(7.00)
    # The existing posting blocker remains as an additional safeguard.
    assert "historical duplicate review holds" in result.metrics["Posting Blockers"]


# ---------------------------------------------------------------------------
# Invariants: nothing disappears, everything ties
# ---------------------------------------------------------------------------

def test_every_row_stays_visible_and_the_row_and_dollar_controls_reconcile(mixed_result):
    result = mixed_result
    source_ids = list(result.qb_work[QB_ID])
    ledger = result.qb_dispositions
    assert list(ledger["QBO Row ID"]) == source_ids                         # every source row, in order

    # Every QuickBooks row appears exactly once on the reconciliation sheet's rows.
    primary_rows = [
        row for row in result.paired_rows if row["QB Index"] is not None and row["QB Record Scope"] == "Primary"
    ]
    assert sorted(row["QB Index"] for row in primary_rows) == sorted(result.qb_work.index)
    sections = {row["Section"] for row in primary_rows}
    assert {"01 Matched", "04 Duplicate QuickBooks", "02 Unmatched QuickBooks", "11 Review Hold QuickBooks"} <= sections

    # Aggregate-match components all remain visible under one group ID.
    assert list(ledger.loc[ledger["Aggregate Group ID"] == "G-001", "QBO Row ID"]) == ["QB-3", "QB-4"]

    # Original rows and dollars = the four dispositions, and the JE is only TRUE_UNMATCHED.
    counts = ledger["Final Disposition"].value_counts()
    assert counts.sum() == result.metrics["QuickBooks Rows"] == len(source_ids)
    by_amount = ledger.groupby("Final Disposition")["Amount"].sum().round(2)
    assert by_amount.sum() == pytest.approx(result.metrics["QuickBooks Source Total"])
    assert result.metrics["Proposed JE Amount"] == pytest.approx(by_amount.get(DISPOSITION_TRUE_UNMATCHED, 0.0))
    assert result.controls.loc[result.controls["Status"] != "PASS"].empty

    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    detail = wb["Reconciliation Detail"]
    assert sum(1 for row in detail.iter_rows(min_row=5) if any(c.value is not None for c in row[:4])) >= len(source_ids)


# ---------------------------------------------------------------------------
# Duplicate confirmation rests only on stable transaction fields
# ---------------------------------------------------------------------------

def test_report_and_export_metadata_never_blocks_confirmation(qb_mapping, inf_mapping, make_metadata):
    """Source row number, report sequence, import timestamp, generated IDs, and any
    informational column are not part of the fingerprint: two records that are the
    same transaction stay confirmed copies however the export labels them."""
    rows = [
        _sales("Acme", "2026-01-05", **{"Export Row #": 7, "Report Seq": 1, "Imported At": "09/01 10:00", "GUID": "a"}),
        _sales("Acme", "2026-01-05", **{"Export Row #": 8, "Report Seq": 2, "Imported At": "09/02 11:30", "GUID": "b"}),
    ]
    result = _reconcile(rows, [_inf("PO-X", "INV-X", 1.00)], _stable(qb_mapping), inf_mapping, make_metadata)
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]


def test_adding_an_informational_column_cannot_change_duplicate_behavior(qb_mapping, inf_mapping, make_metadata):
    def run(extra_a, extra_b, first, second):
        rows = [_sales(first[0], first[1], **extra_a), _sales(second[0], second[1], **extra_b)]
        return _by_disposition(
            _reconcile(rows, [_inf("PO-X", "INV-X", 1.00)], _stable(qb_mapping), inf_mapping, make_metadata),
            DISPOSITION_DUPLICATE_EXCLUDED,
        )

    same, different = ("Acme", "2026-01-05"), ("Beta", "2026-01-09")
    assert run({}, {}, same, same) == run({"Memo": "x"}, {"Memo": "y"}, same, same) == ["QB-2"]
    assert run({}, {}, same, different) == run({"Memo": "x"}, {"Memo": "x"}, same, different) == []


def test_formatting_only_differences_in_stable_fields_still_confirm(qb_mapping, inf_mapping, make_metadata):
    rows = [_sales("Acme  Corp", "1/5/2026"), _sales("ACME CORP", "2026-01-05")]
    result = _reconcile(rows, [_inf("PO-X", "INV-X", 1.00)], _stable(qb_mapping), inf_mapping, make_metadata)
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == ["QB-2"]


def test_a_real_difference_in_customer_date_or_quantity_prevents_confirmation(
    qb_mapping, inf_mapping, make_metadata,
):
    mapping = _stable(qb_mapping)
    for change in ({"Customer": "Beta"}, {"Date": "2026-02-01"}, {"Qty": 2}):
        second = {**_sales("Acme", "2026-01-05"), **change}
        result = _reconcile(
            [_sales("Acme", "2026-01-05"), second], [_inf("PO-X", "INV-X", 1.00)],
            mapping, inf_mapping, make_metadata,
        )
        assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == [], change
        assert _ledger(result).loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE", change


def test_with_no_native_id_and_no_stable_field_mapped_identity_is_not_established(
    qb_mapping, inf_mapping, make_metadata,
):
    """Nothing but the key to go on: the safe, conservative treatment applies."""
    bare = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "quantity": None, "period": None, "product": None}
    rows = [{"PO": "PO1", "Invoice": "INV1", "Amount": 100.00}, {"PO": "PO1", "Invoice": "INV1", "Amount": 100.00}]
    inf = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "period": "Period"}
    result = _reconcile(rows, [_inf("PO-X", "INV-X", 1.00)], bare, inf, make_metadata)
    assert _by_disposition(result, DISPOSITION_DUPLICATE_EXCLUDED) == []
    assert _ledger(result).loc["QB-2", "Reason Code"] == "REVIEW_HOLD_POTENTIAL_DUPLICATE"


# ---------------------------------------------------------------------------
# One standard of evidence across QuickBooks, Infinium, and the historical files
# ---------------------------------------------------------------------------

THIN_INF = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "period": "Period"}


def _thin_inf(po, invoice, amount, **extra):
    return {"PO": po, "Invoice": invoice, "Amount": amount, "Period": "1", **extra}


def test_infinium_copies_need_the_same_sufficient_evidence(qb_mapping, inf_mapping, make_metadata):
    qb_rows = [_qb("PO1", "INV1", 100.00)]
    # Enough line-identity evidence (customer + date identical): a confirmed copy, excluded.
    rich = _reconcile(qb_rows, [_inf("PO1", "INV1", 100.00), _inf("PO1", "INV1", 100.00)],
                      qb_mapping, inf_mapping, make_metadata)
    assert rich.duplicate_inf_rows == [1] and rich.duplicate_review_hold_inf_rows == []
    report = rich.infinium_duplicate_analysis.set_index("Source Row ID")
    assert report.loc["INF-2", "Confirmation Basis"] == "Stable-field fingerprint (Customer, Date)"

    # Nothing but the key and the period (context, not identity): NOT excluded -- held.
    thin = _reconcile(qb_rows, [_thin_inf("PO1", "INV1", 100.00), _thin_inf("PO1", "INV1", 100.00)],
                      qb_mapping, THIN_INF, make_metadata)
    assert thin.duplicate_inf_rows == []
    assert thin.duplicate_review_hold_inf_rows == [1]          # the sibling of the matched row
    assert len(thin.matches) == 1
    reason = thin.infinium_duplicate_analysis.set_index("Source Row ID").loc["INF-2", "Potential Duplicate Reason"]
    assert "not sufficient" in reason
    assert thin.metrics["Control Status"] == "PASS"
    # Both Infinium rows stay in the audit trail.
    sections = sorted(row["Section"] for row in thin.paired_rows if row["Infinium Index"] is not None)
    assert sections == ["01 Matched", "07 Duplicate Review Hold Infinium"]


def test_an_infinium_line_level_id_follows_the_same_contradiction_safeguard(qb_mapping, make_metadata):
    mapping = {**THIN_INF, "line_id": "LineID"}
    qb_rows = [_qb("PO1", "INV1", 100.00)]
    copy = _reconcile(
        qb_rows,
        [_thin_inf("PO1", "INV1", 100.00, LineID="L-1"), _thin_inf("PO1", "INV1", 100.00, LineID="L-1")],
        qb_mapping, mapping, make_metadata,
    )
    assert copy.duplicate_inf_rows == [1]
    assert copy.infinium_duplicate_analysis.set_index("Source Row ID").loc["INF-2", "Confirmation Basis"] == "Line-level source ID"

    both = {**mapping, "customer": "Customer"}
    invoice_level = _reconcile(
        qb_rows,
        [_thin_inf("PO1", "INV1", 100.00, LineID="INV1", Customer="Acme"),
         _thin_inf("PO1", "INV1", 100.00, LineID="INV1", Customer="Beta")],
        qb_mapping, both, make_metadata,
    )
    assert invoice_level.duplicate_inf_rows == []              # shared ID spans different lines: not line-level


def test_historical_copies_are_confirmed_or_withheld_on_the_same_evidence(qb_mapping, inf_mapping, make_metadata):
    qb_rows = [_qb("PO-H", "INV-H", 100.00), _qb("PO-NONE", "INV-NONE", 7.00)]
    inf_rows = [_inf("PO-ZZZ", "INV-ZZZ", 1.00)]

    # Two identical historical rows WITH sufficient evidence: one canonical clears the QuickBooks row.
    confirmed = _reconcile(
        qb_rows, inf_rows, qb_mapping, inf_mapping, make_metadata,
        inf_secondary_raw=pd.DataFrame([_inf("PO-H", "INV-H", 100.00, "12"), _inf("PO-H", "INV-H", 100.00, "12")]),
        inf_secondary_mapping=inf_mapping,
    )
    assert confirmed.duplicate_inf_secondary_rows == [1]
    assert _ledger(confirmed).loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED

    # The same rows with only the key and the period: not confirmed, so BOTH are withheld,
    # nothing is cleared, and the QuickBooks row is held -- never accrued as if absent.
    withheld = _reconcile(
        qb_rows, inf_rows, qb_mapping, inf_mapping, make_metadata,
        inf_secondary_raw=pd.DataFrame([_thin_inf("PO-H", "INV-H", 100.00), _thin_inf("PO-H", "INV-H", 100.00)]),
        inf_secondary_mapping=THIN_INF,
    )
    assert withheld.duplicate_inf_secondary_rows == []
    assert set(withheld.suspected_inf_secondary_rows) == {0, 1}
    ledger = _ledger(withheld)
    assert ledger.loc["QB-1", "Reason Code"] == "REVIEW_HOLD_HISTORICAL_CLEARANCE"
    assert ledger.loc["QB-1", "Related Infinium Row IDs"] == "INF-HIST-1; INF-HIST-2"
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED
    assert withheld.metrics["Proposed JE Amount"] == pytest.approx(7.00)
    assert "historical duplicate review holds" in withheld.metrics["Posting Blockers"]


def test_review_holds_stay_visible_and_every_control_still_ties_across_all_datasets(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reconcile(
        [
            _qb("PO1", "INV1", 100.00),                              # QB-1 matched
            _qb("PO1", "INV1", 100.00, Customer="Beta"),             # QB-2 potential duplicate (held)
            _qb("PO1", "INV1", 100.00),                              # QB-3 confirmed copy of QB-1 (excluded)
            _qb("PO-H", "INV-H", 55.00),                             # QB-4 supported only by withheld historical rows
            _qb("PO-NONE", "INV-NONE", 9.00),                        # QB-5 genuinely missing
        ],
        [
            _thin_inf("PO1", "INV1", 100.00),
            _thin_inf("PO-DUP", "INV-DUP", 30.00), _thin_inf("PO-DUP", "INV-DUP", 30.00),   # unconfirmed copies
            _thin_inf("PO-ZZZ", "INV-ZZZ", 1.00),
        ],
        qb_mapping, THIN_INF, make_metadata,
        inf_secondary_raw=pd.DataFrame([_thin_inf("PO-H", "INV-H", 55.00), _thin_inf("PO-H", "INV-H", 55.00)]),
        inf_secondary_mapping=THIN_INF,
    )
    # 3. only TRUE_UNMATCHED feeds the JE
    ledger = result.qb_dispositions
    fed = ledger.loc[ledger["In Proposed JE"] == "Yes"]
    assert set(fed["Final Disposition"]) == {DISPOSITION_TRUE_UNMATCHED}
    assert result.metrics["Proposed JE Amount"] == pytest.approx(fed["Amount"].sum()) == pytest.approx(9.00)
    # 2. row and dollar controls tie
    assert result.metrics["Control Status"] == "PASS" and result.controls["Status"].eq("PASS").all()
    assert len(ledger) == result.metrics["QuickBooks Rows"] == 5
    assert round(ledger["Amount"].sum(), 2) == pytest.approx(result.metrics["QuickBooks Source Total"])
    # 4. every held / excluded record is still visible in the audit trail
    held = set(ledger.loc[ledger["Final Disposition"] == DISPOSITION_REVIEW_HOLD, "QBO Row ID"])
    assert {"QB-2", "QB-4"} <= held
    duplicate = result.duplicate_analysis.set_index("Source Row ID")
    assert {"QB-1", "QB-2", "QB-3"} <= set(duplicate.index)
    infinium = result.infinium_duplicate_analysis
    assert {"INF-2", "INF-3"} <= set(infinium.loc[infinium["Source Scope"] == "Primary", "Source Row ID"])
    assert result.duplicate_inf_rows == [] and result.duplicate_review_hold_inf_rows == [1, 2]
    assert {"INF-HIST-1", "INF-HIST-2"} <= set(infinium["Source Row ID"])
    primary_qb = [r for r in result.paired_rows if r["QB Index"] is not None and r["QB Record Scope"] == "Primary"]
    assert sorted(r["QB Index"] for r in primary_qb) == sorted(result.qb_work.index)
    primary_inf = [r for r in result.paired_rows if r["Infinium Index"] is not None and r["Infinium Record Scope"] == "Primary"]
    assert sorted(r["Infinium Index"] for r in primary_inf) == sorted(result.inf_work.index)
