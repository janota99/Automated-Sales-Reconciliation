"""Smoke tests for workpapers.py against a real ReconciliationResult.

These exist because matching.py/duplicates.py can be internally correct and
fully unit-tested while still breaking the app: ui_components.py and
workpapers.py each build their own dataframes/column lookups from a
ReconciliationResult, and nothing in test_matching.py or test_duplicates.py
would ever call them. That gap is exactly how a KeyError on
"Reconciliation Status" and a would-be crash on "Source Row IDs" reached a
running Streamlit session undetected. These tests build both real
workbooks end-to-end (with and without duplicates present) so a future
schema change to any ReconciliationResult field gets caught here first.
"""

import io
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from matching import QB_ID, build_reconciliation
from workpapers import (
    _legacy_matched_label,
    build_analytics_workbook,
    build_data_search_dataframe,
    build_legacy_workbook,
    build_primary_workbook,
)

EXPECTED_PRIMARY_SHEETS = [
    "Data Search",
    "Data Search QB Source",
    "Data Search INF Source",
    "Raw Data",
    "Reconciliation Detail",
    "Unresolved Exceptions",
    "Product Aggregate Summary",
]

EXPECTED_LEGACY_SHEETS = [
    "Legacy Reconciliation",
    "Exceptions",
    "Product Aggregate Summary",
]

EXPECTED_ANALYTICS_SHEETS = [
    "Executive Summary",
    "Match Method Summary",
    "Match Register",
    "Detailed Match Ledger",
    "Normalization Detail",
    "Match Assessment",
    "Historical Clearances",
    "Exception Analysis",
    "QuickBooks Duplicates",
    "Infinium Duplicates",
    "Rules and Run Config",
]


def _worksheet_text(ws) -> list[str]:
    return [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]


def _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata):
    qb_rows = [
        {"PO": "PO100", "Invoice": "INV100", "Amount": 100.00, "Qty": 1, "Period": "1"},  # 1:1 match
        {"PO": "PO200", "Invoice": "INV200", "Amount": 50.00, "Qty": 1, "Period": "1"},    # QB duplicate pair
        {"PO": "PO200", "Invoice": "INV200", "Amount": 50.00, "Qty": 1, "Period": "1"},
        {"PO": "PO999", "Invoice": "INV999", "Amount": 15.00, "Qty": 1, "Period": "1"},    # genuine exception
    ]
    inf_rows = [
        {"PO": "PO100", "Invoice": "INV100", "Amount": 100.00, "Period": "1"},
        {"PO": "PO400", "Invoice": "INV400", "Amount": 20.00, "Period": "1"},              # Infinium duplicate pair
        {"PO": "PO400", "Invoice": "INV400", "Amount": 20.00, "Period": "1"},
    ]
    inf_secondary_rows = [
        {"PO": "POHIST", "Invoice": "INVHIST", "Amount": 5.00, "Period": "12"},
    ]
    qb_secondary_rows = [
        {"PO": "POSEC", "Invoice": "INVSEC", "Amount": 9.00, "Qty": 1, "Period": "12"},
    ]
    return build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
        qb_secondary_raw=pd.DataFrame(qb_secondary_rows),
        inf_secondary_raw=pd.DataFrame(inf_secondary_rows),
        qb_secondary_mapping=qb_mapping,
        inf_secondary_mapping=inf_mapping,
    )


def _build_result_without_duplicates(qb_mapping, inf_mapping, make_metadata):
    qb_rows = [
        {"PO": "PO1", "Invoice": "INV1", "Amount": 100.00, "Qty": 1, "Period": "1"},
        {"PO": "PO2", "Invoice": "INV2", "Amount": 15.00, "Qty": 1, "Period": "1"},  # unresolved
    ]
    inf_rows = [
        {"PO": "PO1", "Invoice": "INV1", "Amount": 100.00, "Period": "1"},
    ]
    return build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )


def test_primary_workbook_builds_with_duplicates_present(qb_mapping, inf_mapping, make_metadata):
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    # Each pair is a full-payload strict duplicate: one row is retained as
    # canonical (and, matching nothing else, becomes its own unresolved
    # exception), so only the excess copy is counted as excluded.
    assert result.metrics["Duplicate QuickBooks Rows"] == 1
    assert result.metrics["Duplicate Infinium Rows"] == 1

    workbook_bytes = build_primary_workbook(result)
    assert len(workbook_bytes) > 0

    wb = load_workbook(io.BytesIO(workbook_bytes))
    assert wb.sheetnames == EXPECTED_PRIMARY_SHEETS

    unresolved_text = _worksheet_text(wb["Unresolved Exceptions"])
    assert any("QUICKBOOKS DUPLICATES EXCLUDED FROM JE" in text for text in unresolved_text)
    # The old schema's columns must never resurface in the rendered output.
    assert not any(text == "Reconciliation Status" for text in unresolved_text)
    assert not any(text == "Source Row IDs" for text in unresolved_text)
    # The Infinium exceptions table (its own title band, KPIs, and totals)
    # carries no accrual impact and is already listed in Reconciled Data --
    # it must be gone. QB exception dispositions legitimately mention
    # Infinium (e.g. "No matching Infinium records"), so only the removed
    # section's own headings are checked, not the substring "Infinium".
    assert not any("INFINIUM EXCEPTIONS" in text.upper() for text in unresolved_text)
    assert not any("UNRESOLVED INFINIUM ROWS" in text.upper() for text in unresolved_text)
    assert not any("INFINIUM REVIEW TOTAL" in text.upper() for text in unresolved_text)
    # The duplicates block gets its own reviewer note column, mirroring the
    # QuickBooks exception table's reviewer workflow.
    assert any(text == "Reviewer Note" for text in unresolved_text)


def _fill_matches(cell, hex_color: str) -> bool:
    rgb = cell.fill.fgColor.rgb
    return isinstance(rgb, str) and rgb.endswith(hex_color)


_LEGACY_MATCHED_SECTIONS_FOR_TEST = {
    "01 Matched", "01 Matched - Historical Clearance", "09 Fuzzy Match Review Hold",
}


def test_legacy_workbook_builds_with_duplicates_present(qb_mapping, inf_mapping, make_metadata):
    from config import LEGACY_EXCLUDED_FILL, LEGACY_MATCHED_FILL, LEGACY_METHOD_FILL, LEGACY_REVIEW_FILL

    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    # Sanity check the fixture actually has an Infinium-only exception (the
    # PO400/INV400 duplicate pair, which has no QuickBooks counterpart) --
    # otherwise the QB-only filter below would pass vacuously.
    assert result.metrics["Duplicate Infinium Rows"] == 1
    infinium_only_exceptions = [
        r for r in result.paired_rows
        if r.get("Section") not in _LEGACY_MATCHED_SECTIONS_FOR_TEST and r.get("QB Index") is None
    ]
    assert len(infinium_only_exceptions) >= 1

    workbook_bytes = build_legacy_workbook(result)
    assert len(workbook_bytes) > 0

    wb = load_workbook(io.BytesIO(workbook_bytes))
    assert wb.sheetnames == EXPECTED_LEGACY_SHEETS

    recon_ws = wb["Legacy Reconciliation"]
    recon_text = _worksheet_text(recon_ws)
    assert any("Match Result" in text for text in recon_text)
    assert not any("Match Method" in text or "MATCH METHOD" in text for text in recon_text)
    # The one clean 1:1 match (PO100/INV100/$100) must appear, in a pale
    # green tint, with its method labeled by how it matched -- no
    # confidence tier attached.
    assert any(text == "Unique Match: PO + Invoice + Amount" for text in recon_text)
    # Row 3 is the color legend; row 4 is the header row (unchanged headers).
    header_cells = {cell.value: cell.column for cell in next(recon_ws.iter_rows(min_row=4, max_row=4))}
    method_col = header_cells["Match Result"]
    data_row_cells = list(recon_ws.iter_rows(min_row=5, max_row=5))[0]
    non_method_cells = [c for c in data_row_cells if c.column != method_col]
    assert any(_fill_matches(cell, LEGACY_MATCHED_FILL) for cell in non_method_cells)
    # Body text is regular-weight black -- never bold, never colored.
    for cell in non_method_cells:
        if cell.value is not None:
            assert cell.font.bold is not True
            assert cell.font.color.rgb.endswith("000000")
    # The Match Result data cell is RGB(234,234,234); a clean match isn't bolded.
    method_cell = next(c for c in data_row_cells if c.column == method_col)
    assert _fill_matches(method_cell, LEGACY_METHOD_FILL)
    assert method_cell.font.bold is not True
    # Row 2 (the caption band) stays fixed at height 30; row 1 is pinned to 27.
    assert recon_ws.row_dimensions[2].height == 30
    assert recon_ws.row_dimensions[1].height == 27

    # Every QB row (matched or not) is on Legacy Reconciliation, plus the
    # qualifying Infinium-only rows (any-period duplicate, current-period
    # unmatched) -- not just confirmed matches.
    qb_headers = list(result.qb_raw.columns)
    n_qb = len(qb_headers)
    qualifying_inf_only = [
        r for r in result.paired_rows
        if r.get("QB Index") is None and r.get("Infinium Index") is not None
        and r.get("Section") in {"03 Unmatched Infinium", "05 Duplicate Infinium"}
    ]
    all_qb_rows = [r for r in result.paired_rows if r.get("QB Index") is not None]
    expected_recon_row_count = len(all_qb_rows) + len(qualifying_inf_only)
    total_row_number = next(
        row[0].row for row in recon_ws.iter_rows(min_row=5)
        if any(str(cell.value or "").endswith("TOTAL") for cell in row)
    )
    recon_data_rows = sum(
        1 for row in recon_ws.iter_rows(min_row=5, max_row=total_row_number - 1)
        if any(cell.value is not None for cell in row)
    )
    assert recon_data_rows == expected_recon_row_count

    exceptions_ws = wb["Exceptions"]
    exceptions_text = _worksheet_text(exceptions_ws)
    assert any("Fiscal Period" in text for text in exceptions_text)
    assert any("Exception Type" in text for text in exceptions_text)
    assert "QUICKBOOKS SIDE" in " ".join(exceptions_text).upper()
    fill_colors = {
        cell.fill.fgColor.rgb
        for row in exceptions_ws.iter_rows()
        for cell in row
        if isinstance(cell.fill.fgColor.rgb, str)
    }
    assert any(color.endswith(LEGACY_REVIEW_FILL) for color in fill_colors)
    assert any(color.endswith(LEGACY_EXCLUDED_FILL) for color in fill_colors)
    assert exceptions_ws.row_dimensions[1].height == 27

    # Two independent blocks: general QuickBooks exceptions on the left,
    # excluded QuickBooks duplicate copies on the right, separated by a
    # blank column -- no Infinium columns anywhere on this sheet.
    header_row_cells = next(
        row for row in exceptions_ws.iter_rows()
        if any(cell.value == "Exception Type" for cell in row)
    )
    header_row_number = header_row_cells[0].row
    exception_type_cols = [cell.column for cell in header_row_cells if cell.value == "Exception Type"]
    assert len(exception_type_cols) == 2  # one per block
    left_col, right_col = sorted(exception_type_cols)
    # The right block starts n_qb + 4 (trailer: Fiscal Period, Exception Type,
    # Referenced Match Ref., Explanation) + 1 (separator) columns after the
    # left block's own "Exception Type" column.
    assert right_col - left_col == n_qb + 4 + 1

    right_qb_start_col = right_col - 1 - n_qb  # "Exception Type" is trailer offset +1
    general_pos = [
        row[0].value for row in exceptions_ws.iter_rows(
            min_row=header_row_number + 1, min_col=1, max_col=1,
        ) if row[0].value is not None
    ]
    duplicate_pos = [
        row[0].value for row in exceptions_ws.iter_rows(
            min_row=header_row_number + 1, min_col=right_qb_start_col, max_col=right_qb_start_col,
        ) if row[0].value is not None
    ]
    expected_general = [
        r for r in result.paired_rows
        if r.get("Section") not in _LEGACY_MATCHED_SECTIONS_FOR_TEST
        and r.get("QB Index") is not None
        and r.get("Section") != "04 Duplicate QuickBooks"
    ]
    expected_duplicate = [
        r for r in result.paired_rows if r.get("Section") == "04 Duplicate QuickBooks"
    ]
    assert len(general_pos) == len(expected_general)
    assert len(duplicate_pos) == len(expected_duplicate)

    assert "Product Aggregate Summary" in wb.sheetnames


@pytest.mark.parametrize("engine_result, expected", [
    ("PO + Invoice + Amount", "Unique Match: PO + Invoice + Amount"),
    ("PO + Amount", "Unique Match: PO + Amount"),
    ("Invoice + Amount", "Unique Match: Invoice + Amount"),
    ("PO + Invoice + Aggregate Amount (Grouped) [group-level; no line allocation]", "Group Match: PO + Invoice + Amount"),
    ("PO + Aggregate Amount (Grouped) [group-level; no line allocation]", "Group Match: PO + Amount"),
    ("Invoice + Aggregate Amount (Grouped) [group-level; no line allocation]", "Group Match: Invoice + Amount"),
    # A confirmed vendor alias is a PO-field + amount match.
    ("Confirmed Vendor Alias + Amount", "Unique Match: PO + Amount"),
    # A prior-period clearance keeps only the underlying rule, not its prefix.
    ("QuickBooks primary ↔ Infinium prior-period match | Invoice + Amount", "Unique Match: Invoice + Amount"),
])
def test_legacy_match_method_labels_state_only_how_the_match_was_made(engine_result, expected):
    label = _legacy_matched_label(engine_result)
    assert label == expected
    for confidence in ("Strong", "Moderate", "Weak", "Review", "Confirmed"):
        assert confidence not in label


def _assert_outline_only_panel(ws, row, first_col, last_col):
    """A blank side: every cell carries the panel fill, the outer edges are
    drawn, and no border separates its interior cells."""
    from config import LEGACY_NO_PAIR_FILL

    for col in range(first_col, last_col + 1):
        cell = ws.cell(row, col)
        assert _fill_matches(cell, LEGACY_NO_PAIR_FILL)
        assert cell.border.top.style and cell.border.bottom.style
        assert bool(cell.border.left.style) == (col == first_col)
        assert bool(cell.border.right.style) == (col == last_col)


def test_legacy_reconciliation_styling_labels_and_legend(qb_mapping, inf_mapping, make_metadata):
    from config import (
        LEGACY_EXCLUDED_FILL, LEGACY_MATCHED_FILL, LEGACY_METHOD_FILL,
        LEGACY_NO_PAIR_FILL, LEGACY_REVIEW_FILL,
    )

    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]

    # QuickBooks and Infinium share header names in this fixture; setdefault
    # keeps the first (QuickBooks-side) column for each.
    headers: dict = {}
    for cell in ws[4]:
        headers.setdefault(cell.value, cell.column)
    method_col = headers["Match Result"]
    assert headers["Match Ref."] == method_col - 1  # immediately before Match Result
    amount_col = headers["Amount"]
    po_col = headers["PO"]
    inf_start = headers["Referenced Match Ref."] + 1

    rows = list(ws.iter_rows(min_row=5))
    by_label = {}
    for row in rows:
        label = row[method_col - 1].value
        if label and not str(label).endswith("TOTAL"):
            by_label.setdefault(label, []).append(row)

    # Only "how it matched" / why not -- never a confidence tier.
    for label in by_label:
        assert not any(tag in label for tag in ("(Strong)", "(Moderate)", "(Weak)", "(Review)", "(Needs Review)"))

    # Matched: pale green, regular black text, method cell not bold.
    matched = by_label["Unique Match: PO + Invoice + Amount"][0]
    assert _fill_matches(matched[po_col - 1], LEGACY_MATCHED_FILL)
    assert _fill_matches(matched[inf_start - 1], LEGACY_MATCHED_FILL)
    assert _fill_matches(matched[method_col - 1], LEGACY_METHOD_FILL)
    assert matched[method_col - 1].font.bold is not True
    # Descriptive data stays Segoe UI regular; amounts are regular Consolas.
    assert matched[po_col - 1].font.name == "Segoe UI" and matched[po_col - 1].font.bold is not True
    assert matched[amount_col - 1].font.name == "Consolas" and matched[amount_col - 1].font.bold is not True

    # Unmatched QuickBooks row: gold review tint, bold status text, and a
    # gray (not gold) blank Infinium side because nothing is paired there.
    unmatched = by_label["No Matching Infinium Records"][0]
    assert _fill_matches(unmatched[po_col - 1], LEGACY_REVIEW_FILL)
    assert unmatched[method_col - 1].font.bold is True
    assert _fill_matches(unmatched[method_col - 1], LEGACY_METHOD_FILL)
    assert _fill_matches(unmatched[inf_start - 1], LEGACY_NO_PAIR_FILL)
    # The blank side reads as one quiet panel -- filled, outlined only at its
    # edges, NOT merged (merged cells of different sizes stop Excel sorting).
    inf_end = len(ws[4])
    qb_end_col = headers["Match Ref."] - 1  # last QuickBooks column
    unmatched_row = unmatched[0].row
    _assert_outline_only_panel(ws, unmatched_row, inf_start, inf_end)
    matched_row = matched[0].row
    assert matched[po_col - 1].border.right.style is not None       # populated cells keep their gridlines

    # Excluded duplicate copy: pale red, informational (not bolded).
    duplicate = by_label["Duplicate: Excess Copy Excluded"][0]
    assert _fill_matches(duplicate[po_col - 1], LEGACY_EXCLUDED_FILL)
    assert duplicate[method_col - 1].font.bold is not True

    # Compact legend directly under the introductory note.
    legend = [cell.value for cell in ws[3] if cell.value]
    assert legend == ["Reconciled", "Review required", "Excluded duplicate", "No paired record"]
    # Four adjacent chips, each filled with the exact tint it explains.
    for chip, tint in zip(
        ws[3][:4],
        (LEGACY_MATCHED_FILL, LEGACY_REVIEW_FILL, LEGACY_EXCLUDED_FILL, LEGACY_NO_PAIR_FILL),
    ):
        assert _fill_matches(chip, tint)
    assert ws.row_dimensions[3].height == 18

    # Infinium-only row (no QuickBooks record): its blank QuickBooks side is
    # the same near-white gray, while the Infinium side keeps the status tint.
    inf_only = by_label["No Matching QuickBooks Records"][0]
    assert _fill_matches(inf_only[po_col - 1], LEGACY_NO_PAIR_FILL)
    _assert_outline_only_panel(ws, inf_only[0].row, 1, qb_end_col)
    assert _fill_matches(inf_only[inf_start - 1], LEGACY_REVIEW_FILL) or _fill_matches(
        inf_only[inf_start - 1], LEGACY_EXCLUDED_FILL
    )
    assert ws.print_title_rows == "$1:$4"


def test_legacy_review_labels_and_red_row_reference(qb_mapping, inf_mapping, make_metadata):
    """The red-row sentence in the intro note appears only when a red
    (excluded duplicate) row exists; the legend always lists the chip."""
    with_red = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(with_red)))["Legacy Reconciliation"]
    assert "red rows are excluded duplicates" in ws.cell(2, 1).value

    without_red = _build_result_without_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(without_red)))["Legacy Reconciliation"]
    note = ws.cell(2, 1).value
    assert "red" not in note.lower()
    assert "Gold rows require review. See Exceptions for details." in note
    assert "Excluded duplicate" in [cell.value for cell in ws[3]]


def test_legacy_already_matched_duplicate_reads_as_a_review_action(qb_mapping, inf_mapping, make_metadata):
    """An unresolved QuickBooks row whose PO/invoice belongs to an Infinium
    row already matched elsewhere is labeled as a review action."""
    qb_rows = [
        {"PO": "PO-USED", "Invoice": "INV-A", "Amount": 100.00, "Qty": 1, "Period": "1"},
        {"PO": "PO-USED", "Invoice": "INV-B", "Amount": 55.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [{"PO": "PO-USED", "Invoice": "INV-A", "Amount": 100.00, "Period": "1"}]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]
    labels = {cell.value for row in ws.iter_rows(min_row=5) for cell in row if isinstance(cell.value, str)}
    # The label names the exact accepted match the row's PO was already used by.
    assert "Review: PO Already Used by Match M-001" in labels
    assert not any("Another Match" in label or "Value Already Matched" in label for label in labels)


def test_legacy_short_valued_columns_are_wide_enough_for_their_headings(make_metadata, qb_mapping):
    """Infinium's five-digit "Customer No" column is sized to its data, which
    clips its heading (and the autofilter button) to "Customer N"."""
    qb_rows = [{"PO": "PO100", "Invoice": "INV100", "Amount": 100.00, "Qty": 1, "Period": "1"}]
    inf_rows = [{
        "OHAPD": "1", "OHOBDE": "1/09/2026", "OHCO": "WP", "CUNO": 50001, "OHOBNO": "INV100",
        "OHTOTA": 100.00, "OHDESC": "PO 100", "OHPONO": "PO100",
    }]
    inf_mapping = {"po": "OHPONO", "invoice": "OHOBNO", "amount": "OHTOTA", "period": "OHAPD"}
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]
    for cell in ws[4]:
        if cell.value in ("Match Ref.", "Referenced Match Ref."):
            continue  # narrow by design; the long heading wraps (checked below)
        if cell.value:
            assert ws.column_dimensions[get_column_letter(cell.column)].width >= len(str(cell.value)) + 5, cell.value
    widths = {c.value: ws.column_dimensions[get_column_letter(c.column)].width for c in ws[4] if c.value}
    assert 10 <= widths["Match Ref."] <= 14
    assert 14 <= widths["Referenced Match Ref."] <= 18
    customer_no = next(c for c in ws[4] if c.value == "Customer No.")
    assert ws.column_dimensions[get_column_letter(customer_no.column)].width >= 16


def test_legacy_dates_are_real_dates_shown_as_mm_dd_yyyy(qb_mapping, make_metadata):
    """QuickBooks delivers some dates as real dates and others as text, and
    Infinium's are text like "1/09/2026" -- all must end up as real dates
    displayed MM/DD/YYYY, on both legacy sheets."""
    from datetime import datetime

    qb_mapping = dict(qb_mapping)
    qb_rows = [
        {"PO": "PO1", "Invoice": "INV1", "Amount": 10.00, "Qty": 1, "Period": "1",
         "Date": datetime(2026, 3, 3)},
        {"PO": "PO2", "Invoice": "INV2", "Amount": 20.00, "Qty": 1, "Period": "1",
         "Date": "08/17/2026"},
        {"PO": "PO3", "Invoice": "INV3", "Amount": 30.00, "Qty": 1, "Period": "1",
         "Date": "2026-04-10"},
    ]
    inf_rows = [{"PO": "PO1", "Invoice": "INV1", "Amount": 10.00, "Period": "1", "Date": "1/09/2026"}]
    inf_mapping = {"po": "PO", "invoice": "Invoice", "amount": "Amount", "period": "Period"}
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    wb = load_workbook(io.BytesIO(build_legacy_workbook(result)))
    ws = wb["Legacy Reconciliation"]
    date_cols = [c.column for c in ws[4] if c.value == "Date"]
    assert len(date_cols) == 2  # QuickBooks and Infinium
    checked = 0
    for row in ws.iter_rows(min_row=5):
        for col in date_cols:
            cell = row[col - 1]
            if cell.value is None or str(cell.value).endswith("TOTAL"):
                continue
            assert isinstance(cell.value, datetime), f"{cell.coordinate} is {cell.value!r}"
            assert cell.number_format == "mm/dd/yyyy"
            checked += 1
    assert checked >= 4
    exceptions_ws = wb["Exceptions"]
    exception_date_cells = [
        cell for row in exceptions_ws.iter_rows() for cell in row
        if isinstance(cell.value, datetime)
    ]
    assert exception_date_cells
    assert all(cell.number_format == "mm/dd/yyyy" for cell in exception_date_cells)


def test_legacy_intro_note_reports_percent_reconciled_and_us_style_timestamp(qb_mapping, inf_mapping, make_metadata):
    import re

    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]
    note = ws.cell(2, 1).value
    assert re.match(
        r"^\d[\d,]* of \d[\d,]* QuickBooks records reconciled \(\d+\.\d%\)\. "
        r"Gold rows require review; red rows are excluded duplicates\. See Exceptions for details\. "
        r"Generated \d{2}/\d{2}/\d{4} \d{1,2}:\d{2} (AM|PM) [A-Z]{3,4}\.$",
        note,
    ), note


def test_legacy_workbook_suppresses_number_stored_as_text_warnings(qb_mapping, inf_mapping, make_metadata):
    """Invoice/PO/customer numbers are identifiers, correctly stored as
    text -- the legacy workbook must tell Excel not to flag them with green
    triangles, via each sheet's <ignoredErrors> (placed where the schema
    requires it). The primary workpaper does the same; it has a table part,
    so the element must land before <tableParts>."""
    import zipfile

    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    archive = zipfile.ZipFile(io.BytesIO(build_legacy_workbook(result)))
    sheet_parts = [n for n in archive.namelist() if n.startswith("xl/worksheets/sheet")]
    assert len(sheet_parts) == 3
    for name in sheet_parts:
        xml = archive.read(name).decode("utf-8")
        assert xml.count('numberStoredAsText="1"') == 1, name
        # Schema order: <ignoredErrors> follows pageSetup/headerFooter and
        # precedes the trailing drawing/tableParts, or Excel repairs the file.
        assert xml.index("<ignoredErrors>") > xml.index("</sheetData>")
        assert xml.index("<ignoredErrors>") < xml.index("</worksheet>")
        for later in ("<drawing", "<legacyDrawing", "<tableParts", "<extLst"):
            if later in xml:
                assert xml.index("<ignoredErrors>") < xml.index(later)
    # The archive is still a valid workbook openpyxl can read back.
    assert load_workbook(io.BytesIO(build_legacy_workbook(result))).sheetnames == EXPECTED_LEGACY_SHEETS

    primary = zipfile.ZipFile(io.BytesIO(build_primary_workbook(result)))
    primary_parts = [n for n in primary.namelist() if n.startswith("xl/worksheets/sheet")]
    assert len(primary_parts) == len(EXPECTED_PRIMARY_SHEETS)
    saw_table_sheet = False
    for name in primary_parts:
        xml = primary.read(name).decode("utf-8")
        assert xml.count('numberStoredAsText="1"') == 1, name
        assert xml.index("<ignoredErrors>") > xml.index("</sheetData>")
        if "<tableParts" in xml:
            saw_table_sheet = True
            assert xml.index("<ignoredErrors>") < xml.index("<tableParts")
    assert saw_table_sheet
    assert load_workbook(io.BytesIO(build_primary_workbook(result))).sheetnames == EXPECTED_PRIMARY_SHEETS


def test_legacy_reconciliation_normalizes_infinium_column_names(qb_mapping, make_metadata):
    qb_rows = [{"PO": "PO100", "Invoice": "INV100", "Amount": 100.00, "Qty": 1, "Period": "1"}]
    inf_rows = [{
        "OHAPD": "1", "OHOBDE": "1/09/2026", "OHCO": "WP", "CUNO": 50001, "OHOBNO": "INV100",
        "OHTOTA": 100.00, "OHDESC": "PO 100", "OHPONO": "PO100",
    }]
    inf_mapping = {"po": "OHPONO", "invoice": "OHOBNO", "amount": "OHTOTA", "period": "OHAPD"}
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]
    header_values = [cell.value for cell in ws[4]]
    for raw in ("OHAPD", "OHOBDE", "OHCO", "CUNO", "OHOBNO", "OHTOTA", "OHDESC", "OHPONO"):
        assert raw not in header_values
    inf_headers = header_values[header_values.index("Referenced Match Ref.") + 1:]
    assert inf_headers[:8] == [
        "Period", "Date", "Type", "Customer No.", "Invoice No.", "Amount", "Description", "PO No.",
    ]
    # The underlying values still come through under the friendly names.
    data_values = [cell.value for cell in ws[5]]
    assert 50001 in data_values and "INV100" in data_values


def test_legacy_exceptions_bolds_exception_type_only_for_open_items(qb_mapping, inf_mapping, make_metadata):
    from config import LEGACY_EXCLUDED_FILL, LEGACY_REVIEW_FILL

    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Exceptions"]
    header_cells = next(row for row in ws.iter_rows() if any(c.value == "Exception Type" for c in row))
    header_row = header_cells[0].row
    left_col, right_col = sorted(c.column for c in header_cells if c.value == "Exception Type")

    open_item = ws.cell(header_row + 1, left_col)
    assert open_item.value is not None
    assert open_item.font.bold is True and _fill_matches(open_item, LEGACY_REVIEW_FILL)
    duplicate_copy = ws.cell(header_row + 1, right_col)
    assert duplicate_copy.value is not None
    assert duplicate_copy.font.bold is not True and _fill_matches(duplicate_copy, LEGACY_EXCLUDED_FILL)
    # Other body cells in an exception row are regular-weight black.
    body_cell = ws.cell(header_row + 1, 1)
    assert body_cell.font.bold is not True and body_cell.font.color.rgb.endswith("000000")


def test_legacy_workbook_builds_with_no_duplicates(qb_mapping, inf_mapping, make_metadata):
    result = _build_result_without_duplicates(qb_mapping, inf_mapping, make_metadata)
    workbook_bytes = build_legacy_workbook(result)
    assert len(workbook_bytes) > 0
    wb = load_workbook(io.BytesIO(workbook_bytes))
    assert wb.sheetnames == EXPECTED_LEGACY_SHEETS


def test_data_search_sheet_shows_duplicate_pair_and_clean_match(qb_mapping, inf_mapping, make_metadata):
    """The Data Search sheet must let a reviewer look up any PO/Invoice and
    see its real status -- including for a duplicate pair, where the
    retained copy keeps its normal (accrual-relevant) status but must still
    be identifiable as part of the same duplicate group as the excluded copy."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)

    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    assert "Data Search" in wb.sheetnames

    frame = build_data_search_dataframe(result)
    by_qb_id = frame.set_index("QuickBooks Row ID")

    # PO100/INV100 is a clean 1:1 match -- an ordinary Infinium Match with no
    # duplicate involvement.
    matched = by_qb_id.loc["QB-1"]
    assert matched["Status"] == "Infinium Match"
    assert "PO + Invoice + Amount" in matched["Match Type"]
    assert matched["Duplicate Group ID"] == ""

    # QB-2/QB-3 are the PO200/INV200 duplicate pair. QB-2 is the retained
    # canonical row (matches nothing else, so it's Outstanding) and QB-3 is
    # the excluded excess copy (Status = Duplicate) -- both must carry the
    # same Duplicate Group ID.
    canonical = by_qb_id.loc["QB-2"]
    excess = by_qb_id.loc["QB-3"]
    assert canonical["Status"] == "Outstanding (On Accrual List)"
    assert excess["Status"] == "Duplicate"
    assert canonical["Duplicate Group ID"] != ""
    assert canonical["Duplicate Group ID"] == excess["Duplicate Group ID"]


def test_data_search_sheet_is_first_and_live_searchable(qb_mapping, inf_mapping, make_metadata):
    """Data Search must be the first sheet a reviewer sees, back its two
    live-search panels with hidden source sheets (not visible clutter), and
    drive its results with a FILTER() formula keyed off the two input cells
    -- not a static table requiring manual filtering."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))

    assert wb.sheetnames[0] == "Data Search"
    assert wb["Data Search QB Source"].sheet_state == "hidden"
    assert wb["Data Search INF Source"].sheet_state == "hidden"

    ws = wb["Data Search"]
    search_text = _worksheet_text(ws)
    assert "Enter Invoice #" in search_text
    assert "Enter PO #" in search_text
    assert "QUICKBOOKS ITEMS" in search_text
    assert "INFINIUM ITEMS" in search_text

    qb_formula = ws.cell(9, 1).value
    assert isinstance(qb_formula, str) and qb_formula.startswith("=")
    assert "FILTER(" in qb_formula
    # FILTER is a post-2016 dynamic-array "future function" -- Excel always
    # stores it internally with an _xlfn. prefix, and a file written without
    # that prefix (openpyxl's default) is exactly what triggers Excel's "we
    # found a problem with some content" repair prompt on open.
    assert "_xlfn.FILTER(" in qb_formula
    assert "Data Search QB Source" in qb_formula
    assert "SEARCH(" in qb_formula and "ISNUMBER(" in qb_formula
    # No stray unbalanced parens -- a direct regression guard for the class
    # of hand-built-formula bug this codebase has hit before (#NAME?/#REF!
    # errors from malformed table/range references).
    assert qb_formula.count("(") == qb_formula.count(")")

    # Row ID, PO, Invoice, Amount, Status, Match Ref., Match Type, Referenced Match Ref.,
    # Duplicate Group ID, Matched * Row ID, Detail
    inf_col_count = 11
    inf_formula = ws.cell(9, inf_col_count + 2).value
    assert isinstance(inf_formula, str) and inf_formula.startswith("=")
    assert "Data Search INF Source" in inf_formula
    assert inf_formula.count("(") == inf_formula.count(")")


def test_fiscal_period_summary_is_promoted_above_the_exception_table(
    qb_mapping, inf_mapping, make_metadata,
):
    """The fiscal-period breakdown must sit above the raw exception table,
    not buried below it."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]

    def first_row_containing(needle: str) -> int:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and needle in str(cell.value):
                    return cell.row
        raise AssertionError(f"{needle!r} not found in sheet")

    fiscal_row = first_row_containing("QUICKBOOKS EXCEPTIONS BY FISCAL PERIOD")
    exception_table_header_row = first_row_containing("Exception Status")

    assert fiscal_row < exception_table_header_row
    assert fiscal_row <= 10


def test_unresolved_sheet_has_a_status_color_legend(qb_mapping, inf_mapping, make_metadata):
    """A reviewer opening this workpaper cold has no other way to learn what
    each status fill means, so a color key must be present, sit between the
    KPI cards and the fiscal-period section, and cover the sheet's actual
    status colors (duplicate red, urgent red, pending amber, current
    green)."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]

    def first_row_containing(needle: str) -> int:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and needle in str(cell.value):
                    return cell.row
        raise AssertionError(f"{needle!r} not found in sheet")

    legend_row = first_row_containing("Current Period")
    fiscal_row = first_row_containing("QUICKBOOKS EXCEPTIONS BY FISCAL PERIOD")
    kpi_row = first_row_containing("Net Proposed JE Support")

    assert kpi_row < legend_row < fiscal_row
    # The key wraps onto a second row rather than dropping entries.
    legend_texts = [
        str(cell.value) for row in ws.iter_rows(min_row=legend_row, max_row=fiscal_row - 1)
        for cell in row if cell.value
    ]
    assert len(legend_texts) == 6
    # Period classes are named exactly as the fiscal-period summary names them.
    for name in ("Current Period", "Prior Period / pending review", "Urgent Prior Period"):
        assert name in legend_texts
    assert any("duplicate" in text.lower() for text in legend_texts)


def test_reference_matched_amount_variance_is_excluded_from_accrual_and_shown_separately(
    qb_mapping, inf_mapping, make_metadata,
):
    """A QuickBooks row and an Infinium row that are each other's only
    candidate on a shared PO/invoice, but disagree on amount, must never be
    posted at either amount: not counted as an ordinary unresolved
    QuickBooks exception (and therefore never accrued), but still visible
    to a reviewer in a dedicated "likely data entry error" section rather
    than silently disappearing from the workpaper."""
    qb_rows = [
        {"PO": "PO500", "Invoice": "INV500", "Amount": 100.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [
        {"PO": "PO500", "Invoice": "INV500", "Amount": 90.00, "Period": "1"},
    ]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    assert result.unmatched_qb == []
    assert len(result.amount_variance_analysis) == 1
    assert result.metrics["Reference-Matched Amount Variance Rows"] == 1
    assert result.metrics["Unresolved QuickBooks Rows"] == 0

    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]
    sheet_text = _worksheet_text(ws)

    assert any("REFERENCE-MATCHED AMOUNT VARIANCE" in text for text in sheet_text)
    assert any(text == "QB-1" for text in sheet_text)
    assert any(text == "INF-1" for text in sheet_text)
    assert any(isinstance(v, (int, float)) and v == 100.0 for v in [c.value for row in ws.iter_rows() for c in row])
    assert any(isinstance(v, (int, float)) and v == 90.0 for v in [c.value for row in ws.iter_rows() for c in row])
    # It must not also appear as an ordinary exception row/PROPOSED JE total contributor.
    assert not any(text == "Unmatched QuickBooks" for text in sheet_text)


def test_ambiguous_duplicate_is_excluded_from_accrual_and_labeled_separately(
    qb_mapping, inf_mapping, make_metadata,
):
    """A QuickBooks row whose PO/invoice matches more than one still-
    unresolved Infinium row must never be posted as an ordinary exception --
    it is withheld from accrual and must appear, explicitly labeled
    "ambiguous duplicate" (not folded into the confirmed-duplicate or
    amount-variance sections), so a reviewer can research the candidates."""
    qb_rows = [
        {"PO": "PO-AMBIG", "Invoice": "INV-AMBIG", "Amount": 100.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [
        {"PO": "PO-AMBIG", "Invoice": "INV-AMBIG", "Amount": 90.00, "Period": "1"},
        {"PO": "PO-AMBIG", "Invoice": "INV-AMBIG", "Amount": 80.00, "Period": "1"},
    ]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    assert result.unmatched_qb == []
    assert len(result.ambiguous_duplicate_analysis) == 1
    assert result.amount_variance_analysis.empty

    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]
    sheet_text = _worksheet_text(ws)

    assert any("AMBIGUOUS DUPLICATE" in text for text in sheet_text)
    assert any(text == "QB-1" for text in sheet_text)
    # Both Infinium candidates are listed for the reviewer to research.
    assert any("INF-1" in text and "INF-2" in text for text in sheet_text)
    # It must not also appear as an ordinary exception row, nor be folded
    # into the reference-matched amount variance section (no variance rows
    # were generated at all, per the assertion above).
    assert not any(text == "Unmatched QuickBooks" for text in sheet_text)
    assert not any(str(text).startswith("VAR-") for text in sheet_text)


def test_po_reuse_error_stays_in_accrual_and_shows_grouped_detail(
    qb_mapping, inf_mapping, make_metadata,
):
    """A PO reused across 2+ unresolved QuickBooks rows whose grouped total
    disagrees with Infinium's must remain in the exceptions table and
    accrual total (unlike a duplicate, amount variance, or ambiguous
    duplicate, none of which stay), labeled "PO Re-use Error" at the row
    level, with a supplementary grouped-detail section showing the PO,
    both totals, the difference, and both row counts."""
    qb_rows = [
        {"PO": "PO-REUSE1", "Invoice": "INV-A", "Amount": 100.00, "Qty": 1, "Period": "1"},
        {"PO": "PO-REUSE1", "Invoice": "INV-B", "Amount": 50.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [
        {"PO": "PO-REUSE1", "Invoice": "", "Amount": 140.00, "Period": "1"},
    ]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    assert sorted(result.unmatched_qb) == [0, 1]
    assert len(result.po_reuse_errors) == 1

    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]
    sheet_text = _worksheet_text(ws)

    # Both rows stay as ordinary exceptions, labeled PO Re-use Error, and
    # are still part of the JE-support total (they are never pulled into a
    # withheld/excluded section like the other review-hold categories).
    assert sheet_text.count("PO Re-use Error") >= 2
    assert any("PO RE-USE ERROR" in text for text in sheet_text)
    assert any("QB-1" in text and "QB-2" in text for text in sheet_text)
    assert any(isinstance(v, (int, float)) and v == 150.0 for v in [c.value for row in ws.iter_rows() for c in row])
    assert any(isinstance(v, (int, float)) and v == 140.0 for v in [c.value for row in ws.iter_rows() for c in row])


def test_duplicates_and_je_sit_ten_rows_below_the_exceptions_total(
    qb_mapping, inf_mapping, make_metadata,
):
    """Per explicit request: the duplicates block (and the JE section that
    follows it) must be placed exactly 10 standard rows beneath the
    QuickBooks exceptions total row -- clearly separate from the exception
    detail, but still on the same page without excessive scrolling."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]

    def first_row_containing(needle: str) -> int:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and needle in str(cell.value):
                    return cell.row
        raise AssertionError(f"{needle!r} not found in sheet")

    exceptions_total_row = first_row_containing("PROPOSED JE SUPPORT TOTAL")
    duplicates_row = first_row_containing("QUICKBOOKS DUPLICATES EXCLUDED FROM JE")
    je_row = first_row_containing("PROPOSED JOURNAL ENTRY")

    assert duplicates_row == exceptions_total_row + 10
    assert je_row > duplicates_row
    assert duplicates_row > exceptions_total_row


def test_analytics_workbook_builds_with_duplicates_present(qb_mapping, inf_mapping, make_metadata):
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)

    workbook_bytes = build_analytics_workbook(result)
    assert len(workbook_bytes) > 0

    wb = load_workbook(io.BytesIO(workbook_bytes))
    assert wb.sheetnames == EXPECTED_ANALYTICS_SHEETS

    qb_dup_text = _worksheet_text(wb["QuickBooks Duplicates"])
    assert any("QB-2" in text or "QB-3" in text for text in qb_dup_text)

    inf_dup_text = _worksheet_text(wb["Infinium Duplicates"])
    assert any("INF-2" in text or "INF-3" in text for text in inf_dup_text)


def test_unresolved_sheet_duplicate_tables_use_simplified_reviewer_columns(
    qb_mapping, inf_mapping, make_metadata,
):
    """The 'Unresolved Exceptions' sheet's duplicate sub-tables show a small,
    plain-English column set for non-technical reviewers, while the full
    technical schema stays intact on the 'QuickBooks Duplicates' audit sheet
    -- a regression guard for the "display-only, additive" scope of that
    simplification."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    primary_wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    ws = primary_wb["Unresolved Exceptions"]
    unresolved_text = _worksheet_text(ws)

    for header in ("Duplicate ID", "Row ID", "What Was Found", "Reason", "Status"):
        assert header in unresolved_text, f"missing simplified header {header!r}"
    for technical_header in (
        "Screening Stage", "Normalized PO", "Normalized Invoice",
        "Confirmed Copy Set ID", "Duplicate Rule Version", "Payload Confirmed",
    ):
        assert technical_header not in unresolved_text, (
            f"technical header {technical_header!r} leaked into the reviewer-facing sheet"
        )
    assert any("Same PO 200, Invoice INV200" in text for text in unresolved_text)
    assert any("Kept (Original)" in text or "Removed (Duplicate)" in text for text in unresolved_text)

    analytics_wb = load_workbook(io.BytesIO(build_analytics_workbook(result)))
    audit_text = _worksheet_text(analytics_wb["QuickBooks Duplicates"])
    for technical_header in ("Screening Stage", "Normalized PO", "Confirmed Copy Set ID", "Duplicate Rule Version"):
        assert technical_header in audit_text, f"expected full audit column {technical_header!r} to remain"


def test_unresolved_sheet_row_id_links_to_the_correct_reconciled_data_row(
    qb_mapping, inf_mapping, make_metadata,
):
    """Row ID cells in the duplicates and duplicate-review-hold sections
    must link straight to where that same row was originally listed on
    Reconciliation Detail -- and it must be the CORRECT row, not just any link,
    since a wrong target would be worse than no link at all."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    unresolved_ws = wb["Unresolved Exceptions"]
    reconciled_ws = wb["Reconciliation Detail"]

    def reconciled_po_at(row: int) -> Any:
        return reconciled_ws.cell(row, 1).value  # QB block starts at column A

    qb_po_by_id = result.qb_work.set_index(QB_ID)[qb_mapping["po"]].to_dict()

    def qb_po_for(qb_id: str) -> Any:
        return qb_po_by_id.get(qb_id)

    linked_cells = [
        cell for row in unresolved_ws.iter_rows() for cell in row
        if cell.hyperlink is not None
    ]
    assert linked_cells, "expected at least one Row ID hyperlink on Unresolved Exceptions"

    match_col = next(
        cell.column for cell in reconciled_ws[4] if cell.value == "Match Result"
    )

    checked = 0
    round_trips_checked = 0
    for cell in linked_cells:
        qb_id = str(cell.value)
        target = cell.hyperlink.target
        assert target.startswith("#'Reconciliation Detail'!A")
        target_row = int(target.rsplit("A", 1)[1])
        assert reconciled_po_at(target_row) == qb_po_for(qb_id), (
            f"Row ID {qb_id} links to Reconciliation Detail row {target_row}, which doesn't match"
        )
        # Underline signals "clickable" without erasing a meaningful color
        # (e.g. a duplicate-excluded row's Row ID stays red, just underlined),
        # and bold makes every link easy to spot regardless of its color.
        assert cell.font.underline == "single"
        assert cell.font.bold is True
        checked += 1

        # Round trip: Reconciliation Detail's Match Result cell on that same row
        # must link back to exactly the Unresolved Exceptions row we started
        # from -- bidirectional, not just forward.
        reverse_cell = reconciled_ws.cell(target_row, match_col)
        if reverse_cell.hyperlink is not None:
            reverse_target = reverse_cell.hyperlink.target
            assert reverse_target == f"#'Unresolved Exceptions'!A{cell.row}"
            assert reverse_cell.font.bold is True
            assert reverse_cell.font.underline == "single"
            round_trips_checked += 1
    assert checked >= 2  # both the duplicates and the review-hold section have at least one row here
    assert round_trips_checked >= 2  # every linked row above must round-trip back


def test_primary_and_analytics_workbooks_build_with_no_duplicates(qb_mapping, inf_mapping, make_metadata):
    """The empty-duplicate-frame branches must render cleanly too."""
    result = _build_result_without_duplicates(qb_mapping, inf_mapping, make_metadata)
    assert result.metrics["Duplicate QuickBooks Rows"] == 0
    assert result.metrics["Duplicate Infinium Rows"] == 0

    primary_bytes = build_primary_workbook(result)
    analytics_bytes = build_analytics_workbook(result)
    assert len(primary_bytes) > 0
    assert len(analytics_bytes) > 0

    primary_wb = load_workbook(io.BytesIO(primary_bytes))
    analytics_wb = load_workbook(io.BytesIO(analytics_bytes))
    assert primary_wb.sheetnames == EXPECTED_PRIMARY_SHEETS
    assert analytics_wb.sheetnames == EXPECTED_ANALYTICS_SHEETS

    unresolved_text = _worksheet_text(primary_wb["Unresolved Exceptions"])
    assert any(
        "No QuickBooks exact duplicates" in text for text in unresolved_text
    )


def test_unresolved_sheet_formulas_use_valid_structured_references(
    qb_mapping, inf_mapping, make_metadata,
):
    """Regression test for a #NAME? bug: a bare table name (or INDEX(TableName,0,N)
    built from one) is not a valid Excel reference when written as raw formula
    text -- only Excel's own UI auto-converts a typed table name into proper
    Table[Column] structured-reference syntax. Every formula on this sheet must
    use that qualified form outside the table, or the unqualified [Column] form
    inside the table's own totals row."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    ws = wb["Unresolved Exceptions"]

    formula_cells = [
        (cell.coordinate, cell.value)
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert formula_cells, "expected at least one formula cell to check"
    for coordinate, formula in formula_cells:
        assert "INDEX(" not in formula, f"{coordinate} still uses the broken INDEX(TableName,...) form: {formula}"
        assert "ROWS(QuickBooksExceptions)" not in formula, (
            f"{coordinate} references the table by bare name, which Excel cannot resolve: {formula}"
        )

    # The KPI row (outside the table) must use the qualified Table[Column] form.
    kpi_formulas = " ".join(formula for _, formula in formula_cells if "QuickBooksExceptions[" in formula)
    assert "QuickBooksExceptions[Amount]" in kpi_formulas


def test_render_result_duplicate_badge_never_raises(qb_mapping, inf_mapping, make_metadata):
    """Direct regression test for the exact KeyError that reached Streamlit:
    result.duplicate_analysis["Reconciliation Status"] no longer exists, so
    the attention-badge count must be derived from metrics instead."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    metrics = result.metrics
    unresolved_duplicate_groups = int(metrics.get("Duplicate QuickBooks Rows", 0)) + int(
        metrics.get("Duplicate Infinium Rows", 0)
    )
    assert unresolved_duplicate_groups == 2
    assert "Reconciliation Status" not in result.duplicate_analysis.columns


def _build_result_with_review_hold(qb_mapping, inf_mapping, make_metadata):
    """QuickBooks weak-basis duplicates no longer land in Duplicate Review
    Hold under the current policy (see test_full_reconciliation_detects_
    blank_reference_duplicates and test_weak_basis_review_hold_does_not_
    block_control_status in test_matching.py) -- a zero-evidence group is
    now resolved immediately (one canonical exception, the rest excluded)
    rather than deferred to a human. The "Duplicate Review Hold" section on
    Unresolved Exceptions is kept for the structural case of a future
    QuickBooks policy that defers again, so its rendering is tested here
    by patching a real result's duplicate_analysis into that state directly
    rather than by re-deriving it through build_reconciliation, since no
    QuickBooks input can produce it anymore.
    """
    qb_rows = [
        {"PO": "", "Invoice": "INVBLANK", "Amount": 25.00, "Qty": 1, "Period": "1"},
        {"PO": "", "Invoice": "INVBLANK", "Amount": 25.00, "Qty": 1, "Period": "1"},
        {"PO": "PO999", "Invoice": "INV999", "Amount": 15.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [{"PO": "POX", "Invoice": "INVX", "Amount": 1.00, "Period": "1"}]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    held_qb_ids = set(result.qb_work.loc[[0, 1], QB_ID])
    patched = result.duplicate_analysis.copy()
    held_mask = patched["Source Row ID"].isin(held_qb_ids)
    patched.loc[held_mask, "Disposition"] = "Held for review - excluded from proposed JE pending disposition"
    patched.loc[held_mask, "Confidence"] = "Hold"
    patched.loc[held_mask, "Automatically Excluded"] = True
    patched.loc[held_mask, "Canonical Source Row ID"] = ""
    result.duplicate_analysis = patched
    result.duplicate_review_hold_qb_rows = [0, 1]
    result.duplicate_qb_rows = []
    result.metrics["Duplicate Review Hold QuickBooks Rows"] = 2
    result.metrics["Duplicate Review Hold QuickBooks Amount"] = 50.0
    result.metrics["Duplicate QuickBooks Rows"] = 0
    result.metrics["Duplicate QuickBooks Amount"] = 0.0
    result.metrics["Posting Status"] = "REVIEW REQUIRED"
    return result


def test_duplicate_review_hold_section_renders_with_documented_disposition_dropdown(
    qb_mapping, inf_mapping, make_metadata,
):
    """The Duplicate Review Hold section must appear after the duplicates
    section (before the JE), list the held items, and give the reviewer a
    controlled-vocabulary, editable place to record a documented decision
    -- not just a static, uneditable list. (See _build_result_with_review_
    hold for why this state is patched in rather than produced by a real
    QuickBooks scenario under the current policy.)"""
    result = _build_result_with_review_hold(qb_mapping, inf_mapping, make_metadata)
    assert result.metrics["Posting Status"] == "REVIEW REQUIRED"
    assert result.metrics["Duplicate Review Hold QuickBooks Rows"] == 2

    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    ws = wb["Unresolved Exceptions"]

    def first_row_containing(needle: str) -> int:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and needle in str(cell.value):
                    return cell.row
        raise AssertionError(f"{needle!r} not found in sheet")

    duplicates_row = first_row_containing("QUICKBOOKS DUPLICATES EXCLUDED FROM JE")
    review_hold_row = first_row_containing("DUPLICATE REVIEW HOLD")
    je_row = first_row_containing("PROPOSED JOURNAL ENTRY")
    disposition_header_row = first_row_containing("Reviewer Disposition")

    assert duplicates_row < review_hold_row < je_row

    disposition_col = next(
        cell.column for cell in ws[disposition_header_row] if cell.value == "Reviewer Disposition"
    )
    data_rows = range(disposition_header_row + 1, disposition_header_row + 3)
    for row in data_rows:
        cell = ws.cell(row, disposition_col)
        assert cell.value == "Pending Review"
        assert cell.protection.locked is False, "reviewer must be able to edit the disposition cell"

    list_validations = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    disposition_letter = get_column_letter(disposition_col)
    matching_validation = next(
        dv for dv in list_validations
        if f"{disposition_letter}{disposition_header_row + 1}" in str(dv.sqref)
    )
    assert "Pending Review" in matching_validation.formula1
    assert "Confirmed Duplicate - Exclude Permanently" in matching_validation.formula1
    assert "Confirmed Legitimate - Include In JE Manually" in matching_validation.formula1
    assert "Escalated For Investigation" in matching_validation.formula1


def test_no_review_hold_items_renders_empty_section_cleanly(qb_mapping, inf_mapping, make_metadata):
    """When every weak-basis candidate resolves via a match, the Review
    Hold section must still render (with an explanatory 'none' caption)
    rather than crash on an empty frame."""
    qb_rows = [
        {"PO": "", "Invoice": "INV-DUP", "Amount": 25.00, "Qty": 1, "Period": "1"},
        {"PO": "", "Invoice": "INV-DUP", "Amount": 25.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [{"PO": "", "Invoice": "INV-DUP", "Amount": 50.00, "Period": "1"}]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )
    assert result.metrics["Posting Status"] == "READY TO POST"
    workbook_bytes = build_primary_workbook(result)
    ws = load_workbook(io.BytesIO(workbook_bytes))["Unresolved Exceptions"]
    unresolved_text = _worksheet_text(ws)
    assert any("No QuickBooks weak-basis duplicate candidates remain unresolved" in t for t in unresolved_text)


# ---------------------------------------------------------------------------
# Match references in the generated workbooks
# ---------------------------------------------------------------------------

def _reference_result(qb_mapping, inf_mapping, make_metadata):
    """Two one-to-one matches (PO9 listed before PO1 in the file), a grouped
    match, an exception whose PO a match already used, and a duplicate of a
    matched row -- so every reference kind appears at least once."""
    qb_rows = [
        {"PO": "PO9", "Invoice": "INV9", "Amount": 90.00, "Qty": 1, "Period": "1"},
        {"PO": "PO1", "Invoice": "INV1", "Amount": 10.00, "Qty": 1, "Period": "1"},
        {"PO": "PO-G", "Invoice": "INV-A", "Amount": 100.00, "Qty": 1, "Period": "1"},
        {"PO": "PO-G", "Invoice": "INV-B", "Amount": 50.00, "Qty": 1, "Period": "1"},
        {"PO": "PO1", "Invoice": "INV1B", "Amount": 5.00, "Qty": 1, "Period": "1"},
        {"PO": "PO9", "Invoice": "INV9", "Amount": 90.00, "Qty": 1, "Period": "1"},
    ]
    inf_rows = [
        {"PO": "PO1", "Invoice": "INV1", "Amount": 10.00, "Period": "1"},
        {"PO": "PO9", "Invoice": "INV9", "Amount": 90.00, "Period": "1"},
        {"PO": "PO-G", "Invoice": "", "Amount": 150.00, "Period": "1"},
    ]
    return build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping,
        make_metadata(), 2026,
    )


def _sheet_headers(ws, row):
    return [cell.value for cell in ws[row]]


def _link_text(value):
    """The text a Referenced Match Ref. cell displays. The cell is a live link
    (an IFERROR/HYPERLINK/MATCH formula) whose third HYPERLINK argument is
    the shown text."""
    import re

    if isinstance(value, str) and value.startswith("=IFERROR(HYPERLINK("):
        return re.search(r',"([^"]*)"\),"[^"]*"\)$', value).group(1)
    return value


def test_reconciliation_detail_places_match_ref_immediately_before_match_result(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Reconciliation Detail"]
    headers = _sheet_headers(ws, 4)
    ref_index = headers.index("Match Ref.")
    assert headers[ref_index + 1] == "Match Result"          # Match Ref. immediately before
    assert headers[ref_index + 2] == "Referenced Match Ref."
    ref_col = ref_index + 1
    for position, record in enumerate(result.paired_rows):
        row_number = 5 + position
        shown = ws.cell(row_number, ref_col).value
        pointer = _link_text(ws.cell(row_number, ref_col + 2).value)
        if record["Section"] == "01 Matched":
            assert shown == record["Match Ref."] and shown
            assert pointer is None
        else:
            assert shown is None                              # never an accepted-match reference
            assert (pointer or "") == record["Referenced Match Ref."]
    assert any(ws.cell(5 + i, ref_col + 2).value for i in range(len(result.paired_rows)))
    # Narrow but complete: the workbook autofit must not stretch the reference columns.
    assert ws.column_dimensions[get_column_letter(ref_col)].width <= 14
    assert ws.column_dimensions[get_column_letter(ref_col + 2)].width <= 24


def test_legacy_reconciliation_references_stay_with_their_relationship_after_sorting(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]
    headers = _sheet_headers(ws, 4)
    ref_col = headers.index("Match Ref.")
    assert headers[ref_col + 1] == "Match Result"
    invoice_col = headers.index("Invoice")
    # Expected: each QuickBooks invoice's reference, straight from the register.
    qb_ref_by_id = {}
    for entry in result.match_register.to_dict("records"):
        for qb_id in entry["QuickBooks Row IDs"].split("; "):
            qb_ref_by_id[qb_id] = entry["Match Ref."]
    expected = {}
    for index, qb_id in result.qb_work[QB_ID].items():
        if qb_id in qb_ref_by_id:
            expected.setdefault(result.qb_work.at[index, "Invoice"], set()).add(qb_ref_by_id[qb_id])
    seen = {}
    for row in ws.iter_rows(min_row=5):
        invoice, ref = row[invoice_col].value, row[ref_col].value
        if ref and invoice:
            seen.setdefault(invoice, set()).add(ref)
    assert seen == expected
    # The sheet is sorted by PO, not by reference -- so the order differs from
    # M-001, M-002, ... yet every reference is still attached to the right row.
    ref_order = [row[ref_col].value for row in ws.iter_rows(min_row=5) if row[ref_col].value]
    assert ref_order != sorted(ref_order)
    # Exceptions cite the exact match they point at.
    labels = {c.value for row in ws.iter_rows(min_row=5) for c in row if isinstance(c.value, str)}
    assert any(label.startswith("Potential Duplicate of Match M-") for label in labels)
    assert any(label.startswith("Review: PO Already Used by Match M-") for label in labels)


def test_unresolved_exceptions_places_referenced_match_ref_immediately_before_status(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Unresolved Exceptions"]
    header_row = next(row for row in ws.iter_rows() if any(c.value == "Exception Status" for c in row))
    headers = [c.value for c in header_row]
    status_col = headers.index("Exception Status") + 1
    assert headers[status_col - 2] == "Referenced Match Ref."
    pointer_col = status_col - 1
    rows = []
    for number in range(header_row[0].row + 1, header_row[0].row + 1 + len(result.unmatched_qb)):
        rows.append((ws.cell(number, status_col).value, _link_text(ws.cell(number, pointer_col).value)))
    assert len(rows) == 2
    for status, pointer in rows:
        # Every exception here points at a real match, named the same way in both cells.
        assert pointer in {"M-001", "M-002", "G-001"}
        assert pointer in status


def test_every_reference_shown_in_any_workbook_exists_in_the_match_register(
    qb_mapping, inf_mapping, make_metadata,
):
    import re

    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    register = set(result.match_register["Match Ref."])
    assert register == {"M-001", "M-002", "G-001"}
    pattern = re.compile(r"\b[MG]-\d{3,}\b")
    cited = set()
    for builder in (build_primary_workbook, build_legacy_workbook, build_analytics_workbook):
        wb = load_workbook(io.BytesIO(builder(result)))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cited.update(pattern.findall(cell.value))
    # Nothing is cited that is not a real accepted match; the old six-digit
    # run-local ids (M-000012) would also match this pattern and fail here.
    assert cited == register


def test_analytics_match_level_sheets_carry_references_but_summaries_do_not(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_analytics_workbook(result)))

    def headers(name):
        return next(
            [c.value for c in row] for row in wb[name].iter_rows(min_row=1, max_row=6)
            if sum(1 for c in row if c.value) > 3 and row[0].row >= 3
        )

    ledger = headers("Detailed Match Ledger")
    assert ledger.index("Match Ref.") + 1 == ledger.index("Match Result")
    assert "Referenced Match Ref." in ledger
    assessment = headers("Match Assessment")
    assert assessment.index("Match Ref.") + 1 == assessment.index("Match Method")
    assert "Referenced Match Ref." in assessment
    assert "Referenced Match Ref." in headers("QuickBooks Duplicates")
    # High-level summaries have no match-level records to trace.
    for name in ("Match Method Summary", "Exception Analysis", "Executive Summary"):
        assert not any(
            isinstance(c.value, str) and "Match Ref." in c.value
            for row in wb[name].iter_rows() for c in row
        ), name
    primary = load_workbook(io.BytesIO(build_primary_workbook(result)))
    for name in ("Product Aggregate Summary", "Raw Data"):
        assert not any(
            isinstance(c.value, str) and "Match Ref." in c.value
            for row in primary[name].iter_rows() for c in row
        ), name


def test_data_search_panels_show_references_beside_the_match_type(qb_mapping, inf_mapping, make_metadata):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    for source in ("Data Search QB Source", "Data Search INF Source"):
        headers = _sheet_headers(wb[source], 1)
        assert headers.index("Match Ref.") + 1 == headers.index("Match Type")
        assert "Referenced Match Ref." in headers
        refs = {row[headers.index("Match Ref.")].value for row in wb[source].iter_rows(min_row=2)}
        assert {"M-001", "M-002", "G-001"} <= {r for r in refs if r}


def test_reference_feature_changes_no_reconciliation_total(qb_mapping, inf_mapping, make_metadata):
    """Hand-checked totals for the reference scenario -- the traceability
    fields are additive; every amount, count, and journal-entry input is as
    before."""
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    assert result.metrics["Control Status"] == "PASS"
    assert len(result.matches) == 3                                # PO1, PO9, and the PO-G group
    assert result.metrics["Unresolved QuickBooks Rows"] == 2       # PO1/INV1B and the PO9 duplicate copy
    assert result.metrics["Unresolved QuickBooks Amount"] == pytest.approx(5.00 + 90.00)
    assert result.metrics["QuickBooks Source Total"] == pytest.approx(345.00)
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Unresolved Exceptions"]
    assert any("SUM(QuickBooksExceptions[Amount])" in t for t in _worksheet_text(ws))


def test_legacy_reconciliation_can_be_sorted_and_filtered(qb_mapping, inf_mapping, make_metadata):
    """Excel refuses to sort a range holding merged cells of different sizes.
    Every merge on the sheet must sit above the autofilter's header row."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_legacy_workbook(result)))["Legacy Reconciliation"]
    assert ws.auto_filter.ref
    first_data_row = 5
    assert ws.auto_filter.ref.startswith("A4:")
    assert all(r.max_row < first_data_row - 1 for r in ws.merged_cells.ranges)
    assert not any(
        cell.value is not None and type(cell).__name__ == "MergedCell"
        for row in ws.iter_rows(min_row=first_data_row) for cell in row
    )


def test_analytics_match_register_lists_every_relationship_with_tying_amounts(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_analytics_workbook(result)))
    assert wb.sheetnames.index("Match Register") == wb.sheetnames.index("Match Method Summary") + 1
    ws = wb["Match Register"]
    header_row = next(row for row in ws.iter_rows(max_row=8) if any(c.value == "Match Ref." for c in row))
    headers = [c.value for c in header_row]
    assert "Relationship Key" not in headers                       # no internal run-local ids
    records = [
        dict(zip(headers, [c.value for c in row]))
        for row in ws.iter_rows(min_row=header_row[0].row + 1) if row[0].value
    ]
    assert [r["Match Ref."] for r in records] == list(result.match_register["Match Ref."])
    assert {r["Match Ref."] for r in records} == {"M-001", "M-002", "G-001"}
    group = next(r for r in records if r["Match Ref."] == "G-001")
    assert group["Match Type"] == "Grouped"
    assert group["QuickBooks Row Count"] == 2 and group["Infinium Row Count"] == 1
    assert group["QuickBooks Amount"] == pytest.approx(150.00)
    assert group["Infinium Amount"] == pytest.approx(150.00)
    assert group["Amount Difference"] == pytest.approx(0.0)
    assert all(r["Match Type"] in ("One-to-One", "Grouped") for r in records)


# ---------------------------------------------------------------------------
# Modern main workbook: fiscal-period titles, Reconciliation Detail, Unresolved
# Exceptions layout, and match-reference navigation
# ---------------------------------------------------------------------------

INFINIUM_CODE_MAPPING = {"po": "OHPONO", "invoice": "OHOBNO", "amount": "OHTOTA", "period": "OHAPD"}


def _infinium_code_row(po, invoice, amount, period="1"):
    return {
        "OHAPD": period, "OHOBDE": "1/09/2026", "OHCO": "WP", "CUNO": 50001, "OHOBNO": invoice,
        "OHTOTA": amount, "OHDESC": "Sales order", "OHPONO": po,
    }


def _clearance_result(qb_mapping, make_metadata):
    """A QuickBooks record cleared by a prior-period Infinium record, beside an
    ordinary current-period match and an unresolved exception."""
    qb_rows = [
        {"PO": "PO1", "Invoice": "INV1", "Amount": 100.00, "Qty": 1, "Period": "1"},
        {"PO": "PO-H", "Invoice": "INV-H", "Amount": 25.00, "Qty": 1, "Period": "1"},
        {"PO": "PO-X", "Invoice": "INV-X", "Amount": 7.00, "Qty": 1, "Period": "1"},
    ]
    return build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame([_infinium_code_row("PO1", "INV1", 100.00)]),
        qb_mapping, INFINIUM_CODE_MAPPING, make_metadata(), 2026,
        inf_secondary_raw=pd.DataFrame([_infinium_code_row("PO-H", "INV-H", 25.00, "12")]),
        inf_secondary_mapping=INFINIUM_CODE_MAPPING,
    )


def _detail_sheet(result):
    return load_workbook(io.BytesIO(build_primary_workbook(result)))["Reconciliation Detail"]


def test_main_sheet_titles_lead_with_the_fiscal_period(qb_mapping, inf_mapping, make_metadata):
    """Two runs of the same workbook open side by side must be told apart at a
    glance: the three QuickBooks titles start with the run's fiscal period."""
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    result_p5 = build_reconciliation(
        pd.DataFrame([{"PO": "P", "Invoice": "I", "Amount": 1.0, "Qty": 1, "Period": "5"}]),
        pd.DataFrame([{"PO": "P", "Invoice": "I", "Amount": 1.0, "Period": "5"}]),
        qb_mapping, inf_mapping, make_metadata(fiscal_period=5), 2026,
    )
    wb = load_workbook(io.BytesIO(build_primary_workbook(result_p5)))
    prefix = "FISCAL PERIOD 05 - 2026"
    assert wb["Raw Data"]["A1"].value == f"{prefix} | QUICKBOOKS | RAW TRANSACTION DETAIL"
    assert wb["Reconciliation Detail"]["A1"].value == f"{prefix} | QUICKBOOKS | RECONCILIATION DETAIL"
    assert wb["Unresolved Exceptions"]["A1"].value == f"{prefix} | QUICKBOOKS EXCEPTIONS | JOURNAL ENTRY SUPPORT"
    # The other run says period 01 -- the titles differ, which is the point.
    other = load_workbook(io.BytesIO(build_primary_workbook(result)))
    assert other["Reconciliation Detail"]["A1"].value.startswith("FISCAL PERIOD 01 - 2026 | ")
    # The Infinium side keeps its own heading.
    assert wb["Raw Data"]["A1"].value != wb["Raw Data"].cell(1, 12).value


def test_main_sheet_titles_when_no_fiscal_period_was_selected(qb_mapping, inf_mapping, make_metadata):
    result = build_reconciliation(
        pd.DataFrame([{"PO": "P", "Invoice": "I", "Amount": 1.0, "Qty": 1, "Period": "5"}]),
        pd.DataFrame([{"PO": "P", "Invoice": "I", "Amount": 1.0, "Period": "5"}]),
        qb_mapping, inf_mapping, make_metadata(fiscal_period=None), 2026,
    )
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    assert wb["Reconciliation Detail"]["A1"].value == (
        "FISCAL PERIOD NOT SELECTED - 2026 | QUICKBOOKS | RECONCILIATION DETAIL"
    )


def test_reconciliation_detail_is_named_for_what_it_contains(qb_mapping, inf_mapping, make_metadata):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    assert "Reconciled Data" not in wb.sheetnames
    ws = wb["Reconciliation Detail"]
    assert ws["A1"].value.endswith("| QUICKBOOKS | RECONCILIATION DETAIL")
    infinium_title = next(c.value for c in ws[1] if c.value and str(c.value).startswith("INFINIUM"))
    assert infinium_title == "INFINIUM | RECONCILIATION DETAIL"
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "RECONCILED TOTAL" not in text and "QUICKBOOKS | RECONCILED" not in text


def test_reconciliation_detail_control_strip_ties_to_the_rows_beneath_it(
    qb_mapping, inf_mapping, make_metadata,
):
    from workpapers import qb_record_outcomes
    from utils import format_currency

    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = _detail_sheet(result)
    strip = ws["A3"].value
    outcomes = qb_record_outcomes(result)

    # Every primary QuickBooks record lands in exactly one outcome, and the
    # counts agree with the engine's own metrics.
    assert outcomes["records"] == result.metrics["QuickBooks Rows"]
    assert outcomes["excluded"] == result.metrics["Duplicate QuickBooks Rows"]
    primary_ids = set(result.qb_work[QB_ID])
    register_ids = {
        qb_id for entry in result.match_register["QuickBooks Row IDs"] for qb_id in entry.split("; ")
    }
    assert outcomes["reconciled"] == len(register_ids & primary_ids)
    assert outcomes["reconciled"] + outcomes["unresolved"] + outcomes["excluded"] == outcomes["records"]

    rate = outcomes["reconciled"] / outcomes["records"] * 100
    parts = [part.strip() for part in strip.split("|")]
    noun = "record" if outcomes["records"] == 1 else "records"
    duplicates = "duplicate" if outcomes["excluded"] == 1 else "duplicates"
    assert parts == [
        f"{outcomes['records']:,} QuickBooks {noun}",
        f"{outcomes['reconciled']:,} reconciled",
        f"{rate:.1f}%",
        f"{outcomes['unresolved']:,} unresolved",
        f"{outcomes['excluded']:,} {duplicates} excluded",
        f"JE support: {format_currency(result.metrics['Unresolved QuickBooks Amount'])}",
    ]
    # The strip is a single narrow row above the header, with the control
    # result beside it in green for a PASS.
    assert any(r.min_row == 3 and r.max_row == 3 and r.min_col == 1 for r in ws.merged_cells.ranges)
    assert ws.row_dimensions[3].height <= 24
    control = next(c for c in ws[3] if c.value and str(c.value).startswith("Control:"))
    assert control.value == f"Control: {result.metrics['Control Status']}"
    assert control.fill.fgColor.rgb.endswith("E8F3EC")   # GREEN_LIGHT: control passed
    assert [c.value for c in ws[4]][:2] == ["PO", "Invoice"]
    assert ws.cell(5, 1).value is not None or ws.cell(5, 2).value is not None
    assert ws.freeze_panes.endswith("5")                  # strip and header stay in view
    assert ws.print_title_rows == "$1:$4"


def test_reconciliation_detail_strip_uses_singular_nouns():
    from workpapers import _control_strip_text

    class Stub:
        metrics = {"Unresolved QuickBooks Amount": 0.0}
        paired_rows = [{"QB Index": 0, "QB Record Scope": "Primary", "Section": "04 Duplicate QuickBooks"}]

    text = _control_strip_text(Stub())
    assert text.startswith("1 QuickBooks record   |")
    assert "1 duplicate excluded" in text


def test_reconciliation_detail_shows_friendly_infinium_headers_and_keeps_the_codes(
    qb_mapping, make_metadata,
):
    result = _clearance_result(qb_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    ws = wb["Reconciliation Detail"]
    headers = [c.value for c in ws[4]]
    inf_headers = headers[headers.index("Referenced Match Ref.") + 1:]
    assert inf_headers[:8] == [
        "Period", "Date", "Type", "Customer No.", "Invoice No.", "Amount", "Description", "PO No.",
    ]
    for code in ("OHAPD", "OHOBDE", "OHCO", "CUNO", "OHOBNO", "OHTOTA", "OHDESC", "OHPONO"):
        assert code not in headers
    # The code stays one hover away, on the renamed header itself.
    comments = {c.value: c.comment.text for c in ws[4] if c.comment}
    assert comments["Customer No."] == "Infinium field code: CUNO"
    assert comments["Amount"] == "Infinium field code: OHTOTA"
    assert set(comments) == set(inf_headers[:8])
    # ...and the Raw Data sheet keeps every code exactly as uploaded.
    raw_headers = [c.value for c in wb["Raw Data"][3]]
    for code in ("OHAPD", "OHOBDE", "OHCO", "CUNO", "OHOBNO", "OHTOTA", "OHDESC", "OHPONO"):
        assert code in raw_headers
    # The renamed amount column is still an amount: right-aligned header,
    # accounting number format, measured (numeric) font.
    amount_col = len(headers) - len(inf_headers) + inf_headers.index("Amount") + 1
    assert ws.cell(4, amount_col).alignment.horizontal == "right"
    assert "$" in ws.cell(5, amount_col).number_format
    # The data still comes through under the friendly names.
    assert 50001 in [c.value for c in ws[5]]


def test_record_context_is_blank_unless_a_record_was_accepted_from_a_prior_period(
    qb_mapping, make_metadata,
):
    result = _clearance_result(qb_mapping, make_metadata)
    ws = _detail_sheet(result)
    headers = [c.value for c in ws[4]]
    qb_context = headers.index("QuickBooks Record Context") + 1
    inf_context = headers.index("Infinium Record Context") + 1
    qb_values = [ws.cell(r, qb_context).value for r in range(5, 5 + len(result.paired_rows))]
    inf_values = [ws.cell(r, inf_context).value for r in range(5, 5 + len(result.paired_rows))]
    assert not any(qb_values)                       # ordinary primary records: blank
    assert [v for v in inf_values if v] == ["Accepted Prior-Period Record"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Primary QuickBooks Upload" not in text and "Primary Infinium Upload" not in text
    assert "Prior Period Match" not in text


def test_reconciliation_detail_and_exceptions_link_references_to_the_accepted_match(
    qb_mapping, inf_mapping, make_metadata,
):
    """Every populated Referenced Match Ref. is a live link that finds its
    match on Reconciliation Detail by reference (not by a fixed row), so it
    still lands correctly after that sheet is sorted or filtered."""
    import re

    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    detail = wb["Reconciliation Detail"]
    header = [c.value for c in detail[4]]
    ref_letter = get_column_letter(header.index("Match Ref.") + 1)
    ref_column = {detail.cell(r, header.index("Match Ref.") + 1).value for r in range(5, detail.max_row + 1)}
    pattern = re.compile(
        r"=IFERROR\(HYPERLINK\(\"#'Reconciliation Detail'!([A-Z]+)\"&MATCH\(\"([^\"]+)\","
        r"'Reconciliation Detail'!\$\1:\$\1,0\),\"([^\"]*)\"\),\"\3\"\)"
    )

    linked = []
    pointer_col = header.index("Referenced Match Ref.") + 1
    for r in range(5, 5 + len(result.paired_rows)):
        value = detail.cell(r, pointer_col).value
        if value:
            linked.append(value)
    unresolved = wb["Unresolved Exceptions"]
    exceptions_header = next(row for row in unresolved.iter_rows() if any(c.value == "Exception Status" for c in row))
    exceptions_pointer = [c.value for c in exceptions_header].index("Referenced Match Ref.") + 1
    for r in range(exceptions_header[0].row + 1, exceptions_header[0].row + 1 + len(result.unmatched_qb)):
        value = unresolved.cell(r, exceptions_pointer).value
        if value:
            linked.append(value)
    assert len(linked) >= 3          # exceptions on both sheets point at matches here
    for formula in linked:
        found = pattern.fullmatch(formula)
        assert found, formula
        assert found.group(1) == ref_letter        # the Match Ref. column, wherever it sits
        assert found.group(2) == found.group(3)    # looks up the reference it displays
        assert found.group(2) in ref_column        # ...and that match exists on the sheet
    # Match Ref. cells themselves are the destination, not links.
    assert not any(str(v).startswith("=") for v in ref_column if v)
    # Links are visibly links.
    first_link = next(
        detail.cell(r, pointer_col) for r in range(5, 5 + len(result.paired_rows)) if detail.cell(r, pointer_col).value
    )
    assert first_link.font.underline == "single" and first_link.font.bold is True


def test_a_reference_link_naming_several_matches_shows_all_and_links_to_the_first():
    from workpapers import _match_ref_link_formula

    formula = _match_ref_link_formula("M-003; G-001", "K")
    assert 'MATCH("M-003",' in formula and 'HYPERLINK("#\'Reconciliation Detail\'!K"&MATCH' in formula
    assert formula.endswith('"M-003; G-001"),"M-003; G-001")')


def test_reference_link_columns_are_sized_by_the_reference_not_the_formula(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _reference_result(qb_mapping, inf_mapping, make_metadata)
    wb = load_workbook(io.BytesIO(build_primary_workbook(result)))
    unresolved = wb["Unresolved Exceptions"]
    header_row = next(row for row in unresolved.iter_rows() if any(c.value == "Exception Status" for c in row))
    column = [c.value for c in header_row].index("Referenced Match Ref.") + 1
    width = unresolved.column_dimensions[get_column_letter(column)].width
    assert len("Referenced Match Ref.") <= width          # the full heading fits
    assert header_row[column - 1].alignment.horizontal == "center"
    # Sizing reads the reference a link displays, never the formula behind it.
    from excel_styles import _autofit_display_text
    from workpapers import _match_ref_link_formula

    assert _autofit_display_text(_match_ref_link_formula("M-012", "K")) == "M-012"
    assert _autofit_display_text("=SUM(A1:A9)") == "=SUM(A1:A9)"


def test_unresolved_exceptions_kpi_blocks_use_explicit_proposed_entry_labels(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Unresolved Exceptions"]
    labels = [c.value for c in ws[3] if c.value]
    assert labels == [
        "Unresolved Rows", "Proposed Debit Support", "Proposed Credit Support", "Net Proposed JE Support",
    ]
    text = _worksheet_text(ws)
    for old in ("Gross debits", "Credits", "Proposed JE support total", "Unresolved rows"):
        assert old not in text
    # Compact: a slim label row over a slim value row, four cards side by side.
    assert ws.row_dimensions[3].height <= 16 and ws.row_dimensions[4].height <= 24
    assert len([r for r in ws.merged_cells.ranges if r.min_row == 3 and r.max_row == 3]) == 4


def test_unresolved_exceptions_title_to_summary_gap_is_tight(qb_mapping, inf_mapping, make_metadata):
    """The fiscal-period summary sits close beneath the title -- rows are
    sized from a merged caption's full width, not its first column."""
    result = _build_result_with_duplicates(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Unresolved Exceptions"]
    fiscal_row = next(
        c.row for row in ws.iter_rows() for c in row
        if c.value and str(c.value).startswith("QUICKBOOKS EXCEPTIONS BY FISCAL PERIOD")
    )
    above = sum(ws.row_dimensions[r].height or 15 for r in range(1, fiscal_row))
    assert above <= 150, above
    # The summary's own caption is a line or two, not a dozen.
    assert (ws.row_dimensions[fiscal_row + 1].height or 15) <= 45
    # A one-line title band is one designed height everywhere.
    assert ws.row_dimensions[1].height == 27
    raw = load_workbook(io.BytesIO(build_primary_workbook(result)))["Raw Data"]
    assert raw.row_dimensions[1].height == 27


def test_unresolved_exceptions_period_classes_are_named_consistently(qb_mapping, inf_mapping, make_metadata):
    from config import AMBER, GREEN_LIGHT, RED_LIGHT

    qb_rows = [
        {"PO": "PO-CUR", "Invoice": "INV-CUR", "Amount": 10.00, "Qty": 1, "Period": "5"},
        {"PO": "PO-P4", "Invoice": "INV-P4", "Amount": 20.00, "Qty": 1, "Period": "4"},
        {"PO": "PO-P1", "Invoice": "INV-P1", "Amount": 50.00, "Qty": 1, "Period": "1"},
    ]
    result = build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame([{"PO": "POX", "Invoice": "INVX", "Amount": 1.0, "Period": "5"}]),
        qb_mapping, inf_mapping, make_metadata(fiscal_period=5), 2026,
    )
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Unresolved Exceptions"]
    fiscal_header = next(row for row in ws.iter_rows() if any(c.value == "Period Classification" for c in row))
    class_col = [c.value for c in fiscal_header].index("Period Classification") + 1
    classes = {}
    row = fiscal_header[0].row + 1
    while ws.cell(row, 1).value and not str(ws.cell(row, 1).value).startswith("TOTAL"):
        classes[ws.cell(row, class_col).value] = ws.cell(row, class_col)
        # The color sits on the classification cell alone -- not the whole row.
        for col in range(1, class_col + 1):
            if col != class_col:
                assert not ws.cell(row, col).fill.fgColor.rgb.endswith((AMBER, GREEN_LIGHT, RED_LIGHT)), (row, col)
        row += 1
    assert set(classes) == {"Current Period", "Prior Period", "Urgent Prior Period"}
    assert classes["Current Period"].fill.fgColor.rgb.endswith(GREEN_LIGHT)
    assert classes["Prior Period"].fill.fgColor.rgb.endswith(AMBER)
    assert classes["Urgent Prior Period"].fill.fgColor.rgb.endswith(RED_LIGHT)
    # The legend uses those same three names.
    legend = {c.value for r in ws.iter_rows(min_row=5, max_row=fiscal_header[0].row) for c in r if c.value}
    assert "Current Period" in legend and "Urgent Prior Period" in legend
    assert any(text.startswith("Prior Period") for text in legend)
    caption = next(c.value for r in ws.iter_rows() for c in r if c.value and str(c.value).startswith("Selected current"))
    assert "Prior Period exception" in caption and "Urgent Prior Period exception" in caption


def test_unresolved_exceptions_review_hold_tints_only_its_status_cell(qb_mapping, inf_mapping, make_metadata):
    from config import AMBER

    result = _build_result_with_review_hold(qb_mapping, inf_mapping, make_metadata)
    ws = load_workbook(io.BytesIO(build_primary_workbook(result)))["Unresolved Exceptions"]
    title_row = next(
        c.row for row in ws.iter_rows() for c in row if c.value and str(c.value).startswith("DUPLICATE REVIEW HOLD")
    )
    header_row = next(
        row for row in ws.iter_rows(min_row=title_row) if any(c.value == "Reviewer Disposition" for c in row)
    )
    headers = [c.value for c in header_row]
    status_col = headers.index("Status") + 1
    data_row = header_row[0].row + 1
    assert ws.cell(data_row, status_col).value == "Pending Review"
    for col in range(1, len(headers) + 1):
        filled = ws.cell(data_row, col).fill.fgColor.rgb.endswith(AMBER)
        assert filled == (col == status_col), col


def test_row_autofit_measures_a_merged_caption_across_its_full_width():
    """A long caption in a merged band must size from the whole band's width;
    measured against its first column alone it balloons the row."""
    from openpyxl import Workbook

    from workpapers import _autofit_workbook_rows

    wb = Workbook()
    ws = wb.active
    for letter in "ABCDEFGHIJ":
        ws.column_dimensions[letter].width = 20
    ws.merge_cells("A5:J5")
    ws["A5"] = "word " * 60                      # 300 characters across 200 units of width
    ws["A6"] = "=SUMIF(A1:A5,\">0\",A1:A5)+SUMIF(B1:B5,\">0\",B1:B5)+SUMIF(C1:C5,\">0\",C1:C5)"
    _autofit_workbook_rows(wb)
    assert (ws.row_dimensions[5].height or 15) <= 45
    assert ws.row_dimensions[6].height is None    # a formula is not measured as text
